"""Synthetic host tests for the batch-flash safety and persistence contract.

Run with python3 -m unittest discover -s tests -p 'test_batch_flash.py' -v.
All workbook/build fixtures are temporary; serial and process access is mocked.
"""

import contextlib
import io
import json
import os
from pathlib import Path
import sys
import tempfile
import threading
import types
import unittest
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import batch_flash as batch
import openpyxl
from openpyxl.styles import Font, PatternFill

PID = "synthetic-pid-001"


def credential(number):
    return {"uuid": "synthetic-uuid-%03d" % number,
            "key": "synthetic-key-%03d-" % number + "x" * 16}


def tracked(number, status):
    row = credential(number)
    row.update(status=status, stage="complete" if status == "used" else "firmware",
               device_mac="02:00:00:00:00:%02x" % number,
               product_key=PID, port="/dev/synthetic-%d" % number,
               run_id="synthetic-run", updated_at="2026-01-01T00:00:00+00:00")
    if status == "fail":
        row["error"] = "synthetic transport failure"
    return row


def port(path, **changes):
    values = dict(device=path, description="Synthetic USB bridge", hwid="USB test",
                  vid=0x1234, pid=0x5678, serial_number="shared-serial", location="1-2")
    values.update(changes)
    return types.SimpleNamespace(**values)


class FakeTTY(io.StringIO):
    """In-memory stream that claims to be an interactive terminal."""

    def isatty(self):
        return True


class BrokenStream(io.StringIO):
    """Terminal-like stream whose writes always fail."""

    def isatty(self):
        return True

    def write(self, *_args, **_kwargs):
        raise OSError("broken pipe")

    def flush(self):
        pass


def display_assignments(paths):
    return [batch.Assignment(batch.PortCandidate(path),
                             types.SimpleNamespace(index=index + 2), PID)
            for index, path in enumerate(paths)]


class SyntheticCase(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix="batch_test_", dir="/tmp")
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name).resolve()
        self.xlsx = self.root / "synthetic credentials.xlsx"
        self.build = self.root / "build with spaces"
        self.build.mkdir()
        self.manifest_data = {
            "extra_esptool_args": {"chip": "esp32s3", "before": "default_reset",
                                   "after": "hard_reset", "stub": False},
            "partition-table": {"offset": "0x8000", "file": "partition table.bin"},
            "write_flash_args": ["--flash_mode", "dio", "--flash_freq", "80m",
                                 "--flash_size", "16MB"],
            "flash_files": {"0x10000": "app image.bin", "0x0": "bootloader.bin",
                            "0x8000": "partition table.bin", "0xd000": "ota data.bin",
                            "0x900000": "assets image.bin"},
        }
        for name in self.manifest_data["flash_files"].values():
            (self.build / name).write_bytes(b"synthetic image: " + name.encode())
        self.save_manifest()
        self.make_workbook([credential(i) for i in range(1, 5)])
        self.output, self.errors = io.StringIO(), io.StringIO()
        self.stack = contextlib.ExitStack()
        self.addCleanup(self.stack.close)
        self.stack.enter_context(contextlib.redirect_stdout(self.output))
        self.stack.enter_context(contextlib.redirect_stderr(self.errors))
        self.stack.enter_context(mock.patch.dict(os.environ, {"ESPPORT": ""}))
        self.patch(batch, "_install_signal_handlers")
        self.patch(batch, "_validate_idf_tools")
        self.patch(batch, "esptool_subcommand", side_effect=lambda name: name)
        self.commands = self.patch(batch, "run_command",
                                   side_effect=AssertionError("unmocked device command"))
        self.auth = self.patch(batch, "write_verify_identity",
                               side_effect=AssertionError("unmocked auth device access"))
        self.candidates = [batch.PortCandidate("/dev/cu.synthetic-long-path-%03d" % i,
                                              "Synthetic USB bridge", vid=0x1234,
                                              pid=0x5678, serial_number="serial-%d" % i,
                                              location="hub-%d" % i)
                           for i in range(1, 3)]
        self.discovery = self.patch(batch, "discover_candidates",
                                    return_value=(self.candidates, []))

    def patch(self, target, name, **kwargs):
        return self.stack.enter_context(mock.patch.object(target, name, **kwargs))

    def save_manifest(self):
        (self.build / "flasher_args.json").write_text(json.dumps(self.manifest_data))

    def make_workbook(self, rows, headers=None):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Credentials"
        headers = headers or ["uuid", "key", "note"] + list(batch.TRACKING_FIELDS)
        sheet.append(headers)
        for row in rows:
            sheet.append([row.get(field) for field in headers])
        workbook.save(self.xlsx)
        workbook.close()
        self.xlsx.chmod(0o640)

    def load(self, write=False):
        workbook = batch.Workbook(str(self.xlsx)).load(for_write=write)
        self.addCleanup(workbook.wb.close)
        return workbook

    def argv(self, *extra):
        return ["--pid", PID, "--xlsx", str(self.xlsx), "--build-dir", str(self.build),
                "--yes"] + list(extra)

    def args(self, *extra):
        return batch.parse_args(self.argv(*extra))

    def enable_devices(self):
        self.patch(batch, "verify_candidates_frozen")
        self.finalize = self.patch(batch, "finalize_device")
        self.probes = self.patch(batch, "probe_device", side_effect=self.good_probe)
        self.commands.side_effect = None
        self.commands.return_value = types.SimpleNamespace(stdout="", stderr="")
        self.auth.side_effect = self.good_auth

    def good_probe(self, candidate, manifest, baud, timeout, cancel_event):
        index = self.candidates.index(candidate) + 1
        return batch.ProbeResult("ESP32-S3", "02:00:00:00:00:%02x" % index)

    @staticmethod
    def good_auth(path, baud, values, **kwargs):
        for stage in ("auth_write", "auth_read", "auth_verify"):
            kwargs["stage_callback"](stage)
        return True

    def journal_records(self):
        path = Path(str(self.xlsx) + ".batch-journal.jsonl")
        return [json.loads(line) for line in path.read_text().splitlines()]

    def assignments_and_coordinator(self):
        workbook = self.load(write=True)
        assignments = batch.plan_assignments(self.args(), workbook, self.candidates)
        for i, assignment in enumerate(assignments, 1):
            assignment.mac = "02:00:00:00:00:%02x" % i
        journal = batch.Journal(str(self.xlsx))
        self.addCleanup(journal.close)
        coordinator = batch.Coordinator(workbook, journal, "synthetic-run", threading.Event())
        coordinator.register_secrets(assignments)
        coordinator.reserve(assignments, PID)
        return assignments, coordinator


class TestDiscovery(SyntheticCase):
    def enumerate(self, platform, ports):
        with mock.patch.object(batch.sys, "platform", platform), \
             mock.patch.object(batch, "_import_list_ports",
                               return_value=types.SimpleNamespace(comports=lambda **kw: ports)):
            return ORIGINAL_DISCOVER()

    def test_macos_usb_native_and_bridge_filtering_sorted_without_serial_dedupe(self):
        ports = [port("/dev/cu.usbserial-z"), port("/dev/tty.usbserial-z"),
                 port("/dev/cu.usbmodem-a"), port("/dev/cu.Bluetooth-Incoming-Port"),
                 port("/dev/cu.debug-console"), port("/dev/cu.unknown", vid=None,
                                                       hwid="", location=None),
                 port("/dev/ttyS0")]
        selected, excluded = self.enumerate("darwin", ports)
        self.assertEqual([item.path for item in selected],
                         ["/dev/cu.usbmodem-a", "/dev/cu.usbserial-z"])
        self.assertEqual(len(excluded), 5)
        self.assertEqual(selected[0].vid_pid, "1234:5678")
        self.assertEqual(selected[0].description, "Synthetic USB bridge")
        self.assertEqual(selected[0].location, "1-2")
        self.assertEqual(selected[0].serial_number, selected[1].serial_number)

    def test_linux_usb_without_metadata_and_nonusb_with_vid(self):
        selected, excluded = self.enumerate("linux", [
            port("/dev/ttyUSB1", vid=None, pid=None, hwid="", location=None),
            port("/dev/ttyACM0"), port("/dev/ttyS0"), port("/dev/ttyAMA0"),
            port("/dev/rfcomm0"), port("/dev/arbitrary")])
        self.assertEqual([item.path for item in selected], ["/dev/ttyACM0", "/dev/ttyUSB1"])
        self.assertEqual(len(excluded), 4)

    def test_canonical_and_macos_alias_deduplication(self):
        device = self.root / "device"
        device.touch()
        alias = self.root / "alias"
        alias.symlink_to(device)
        selected, dropped = batch._dedupe_candidates([
            batch.PortCandidate(str(device)), batch.PortCandidate(str(alias)),
            batch.PortCandidate("/dev/tty.usb-test"), batch.PortCandidate("/dev/cu.usb-test")])
        self.assertEqual(len(selected), 2)
        self.assertEqual(len(dropped), 2)
        self.assertIn("/dev/cu.usb-test", [item.path for item in selected])

    def test_explicit_paths_keep_argument_order_without_metadata(self):
        paths = [self.root / "z", self.root / "a"]
        for path in paths:
            path.touch()
        with mock.patch.object(batch, "_metadata_map", return_value={}):
            candidates = batch.explicit_candidates([str(path) for path in paths])
        self.assertEqual([c.path for c in candidates], [str(path) for path in paths])
        self.assertTrue(all(c.vid is None for c in candidates))

    def test_explicit_paths_reject_duplicates_aliases_relative_and_missing(self):
        for paths in (["/dev/fake", "/dev/fake"],
                      ["/dev/cu.fake", "/dev/tty.fake"], ["slot=/dev/fake"],
                      [str(self.root / "missing")]):
            with self.subTest(paths=paths), \
                 mock.patch.object(batch, "_metadata_map", return_value={}), \
                 self.assertRaises(batch.BatchError):
                batch.explicit_candidates(paths)

    def test_snapshot_ignores_new_arrivals_and_does_not_mutate_selection(self):
        candidate = self.candidates[0]
        fresh = port(candidate.path, serial_number=candidate.serial_number,
                     location=candidate.location)
        with mock.patch.object(batch, "_metadata_map", return_value={
                candidate.path: fresh, "/dev/new": port("/dev/new")}), \
             mock.patch.object(batch.os.path, "exists", return_value=True):
            snapshot = [candidate]
            batch.verify_candidates_frozen(snapshot)
        self.assertEqual(snapshot, [candidate])

    def test_snapshot_rejects_every_observable_metadata_change(self):
        candidate = self.candidates[0]
        for field, value in (("vid", 1), ("pid", 2), ("serial_number", "other"),
                             ("location", "other")):
            fresh = port(candidate.path, serial_number=candidate.serial_number,
                         location=candidate.location)
            setattr(fresh, field, value)
            with self.subTest(field=field), \
                 mock.patch.object(batch, "_metadata_map", return_value={candidate.path: fresh}), \
                 mock.patch.object(batch.os.path, "exists", return_value=True), \
                 self.assertRaisesRegex(batch.BatchError, field):
                batch.verify_candidates_frozen([candidate])

    def test_snapshot_rejects_disappeared_path_or_metadata_without_substitution(self):
        for exists in (True, False):
            with self.subTest(exists=exists), \
                 mock.patch.object(batch, "_metadata_map", return_value={"/dev/new": port("/dev/new")}), \
                 mock.patch.object(batch.os.path, "exists", return_value=exists), \
                 self.assertRaisesRegex(batch.BatchError, "disappeared"):
                batch.verify_candidates_frozen([self.candidates[0]])

    def test_metadata_less_explicit_snapshot_requires_only_existing_path(self):
        with mock.patch.object(batch, "_metadata_map", return_value={}), \
             mock.patch.object(batch.os.path, "exists", return_value=True):
            batch.verify_candidates_frozen([batch.PortCandidate("/dev/synthetic")])


class TestManifest(SyntheticCase):
    def test_full_image_command_preserves_paths_spaces_offsets_and_settings(self):
        manifest = batch.Manifest(str(self.build)).load()
        command = manifest.write_flash_command("/dev/port with spaces", 460800)
        expected_images = [part for entry in manifest.entries for part in entry]
        self.assertEqual(command[-10:], expected_images)
        self.assertEqual([offset for offset, path in manifest.entries],
                         ["0x0", "0x8000", "0xd000", "0x10000", "0x900000"])
        self.assertEqual(command[:7], [sys.executable, "-m", "esptool", "--chip", "esp32s3",
                                      "--port", "/dev/port with spaces"])
        self.assertIn("--no-stub", command)
        self.assertEqual(command[command.index("--after") + 1], "no_reset")
        self.assertEqual(command[command.index("--before") + 1], "default_reset")
        self.assertEqual(command[command.index("--baud") + 1], "460800")
        self.assertEqual(command[command.index("write_flash") + 1:-10],
                         self.manifest_data["write_flash_args"])
        self.assertTrue(all(Path(path).is_absolute() for _, path in manifest.entries))

    def test_encryption_markers_and_force_flags_rejected(self):
        for value in (True, 1, "true", "1", "yes", [1]):
            with self.subTest(marker=value):
                self.manifest_data["app"] = {"encrypted": value}
                self.save_manifest()
                with self.assertRaisesRegex(batch.BatchError, "encrypted"):
                    batch.Manifest(str(self.build)).load()
        del self.manifest_data["app"]
        for flag in ("--force", "--encrypt", "--encrypt-files=x", "--ignore_flash_enc_efuse"):
            with self.subTest(flag=flag):
                self.manifest_data["write_flash_args"] = [flag]
                self.save_manifest()
                with self.assertRaisesRegex(batch.BatchError, "unsupported"):
                    batch.Manifest(str(self.build)).load()

    def test_false_encryption_markers_accepted(self):
        for value in (False, 0, "false", "0", "no", "", None):
            with self.subTest(value=value):
                self.manifest_data["app"] = {"encrypted": value}
                self.save_manifest()
                batch.Manifest(str(self.build)).load()

    def test_invalid_missing_duplicate_or_escaping_images_rejected(self):
        outside = self.root / "outside.bin"
        outside.write_bytes(b"outside")
        (self.build / "escape.bin").symlink_to(outside)
        for files in ({}, {"123": "bootloader.bin"},
                      {"0x0": "bootloader.bin", "0x00": "app image.bin"},
                      {"0x0": "missing.bin"}, {"0x0": "../outside.bin"},
                      {"0x0": str(outside)}, {"0x0": "escape.bin"}, {"0x0": None}):
            with self.subTest(files=files):
                self.manifest_data["flash_files"] = files
                self.save_manifest()
                with self.assertRaises(batch.BatchError):
                    batch.Manifest(str(self.build)).load()

    def test_fingerprint_detects_same_size_same_mtime_image_edits(self):
        manifest = batch.Manifest(str(self.build)).load()
        path = self.build / "bootloader.bin"
        stat = path.stat()
        path.write_bytes(b"z" * stat.st_size)
        os.utime(path, ns=(stat.st_atime_ns, stat.st_mtime_ns))
        with self.assertRaisesRegex(batch.BatchError, "changed"):
            manifest.verify_fingerprint()

    def test_fingerprint_detects_manifest_edit_or_removed_image(self):
        for kind in ("manifest", "missing"):
            with self.subTest(kind=kind):
                self.save_manifest()
                manifest = batch.Manifest(str(self.build)).load()
                if kind == "manifest":
                    with (self.build / "flasher_args.json").open("a") as handle:
                        handle.write(" ")
                else:
                    (self.build / "bootloader.bin").unlink()
                with self.assertRaises(batch.BatchError):
                    manifest.verify_fingerprint()


class TestWorkbook(SyntheticCase):
    def test_sheet_header_detection_empty_rows_and_explicit_sheet(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Pool"
        sheet.append(["Synthetic credential fixture"])
        sheet.append([" UUID ", "Key", "Note"])
        sheet.append([None, None, "unrelated row"])
        sheet.append(list(credential(1).values()))
        workbook.create_sheet("Notes")["A1"] = "preserve me"
        workbook.save(self.xlsx)
        workbook.close()
        loaded = self.load()
        self.assertEqual(loaded.header_row, 2)
        self.assertEqual(list(loaded.rows), [4])
        self.assertEqual(loaded.sheet_title, "Pool")
        explicit = batch.Workbook(str(self.xlsx), "Pool").load()
        self.addCleanup(explicit.wb.close)
        self.assertEqual(explicit.rows[4].uuid, credential(1)["uuid"])

    def test_missing_duplicate_and_ambiguous_headers_rejected(self):
        for headers in (["uuid", "missing"], ["uuid", "key", " UUID "],
                        ["uuid", "key", "status", " Status "]):
            with self.subTest(headers=headers):
                self.make_workbook([credential(1)], headers)
                with self.assertRaises(batch.BatchError):
                    self.load()
        self.make_workbook([credential(1)])
        workbook = openpyxl.load_workbook(self.xlsx)
        workbook.copy_worksheet(workbook.active)
        workbook.save(self.xlsx)
        workbook.close()
        with self.assertRaisesRegex(batch.BatchError, "multiple worksheets"):
            self.load()
        with self.assertRaisesRegex(batch.BatchError, "duplicate"):
            batch.Workbook(str(self.xlsx), "Credentials").load()
        with self.assertRaisesRegex(batch.BatchError, "worksheet not found"):
            batch.Workbook(str(self.xlsx), "Missing").load()

    def test_partial_numeric_formula_whitespace_and_byte_lengths_rejected(self):
        for field in ("uuid", "key"):
            low, high = batch.LENGTH_LIMITS["auth_key" if field == "key" else field]
            for value in (None, 123, "=REPT(\"x\",32)", " " + "a" * low,
                          "a" * low + " ", "a" * (low - 1), "a" * (high + 1),
                          "\u00e9" * (high // 2 + 1)):
                with self.subTest(field=field, value=value):
                    row = credential(1)
                    row[field] = value
                    self.make_workbook([row])
                    with self.assertRaises(batch.BatchError) as caught:
                        self.load()
                    self.assertIn("row 2", str(caught.exception))
                    self.assertNotIn(credential(1)["key"], str(caught.exception))

    def test_credential_nul_and_exact_utf8_boundaries(self):
        workbook = self.load()
        for field in ("uuid", "key"):
            low, high = batch.LENGTH_LIMITS["auth_key" if field == "key" else field]
            for length in (low, high):
                value = "\u00e9" * (length // 2) + "a" * (length % 2)
                self.assertEqual(workbook._credential_value(value, field, 2), value)
            with self.assertRaisesRegex(batch.BatchError, "NUL"):
                workbook._credential_value("a" * low + "\x00", field, 2)

    def test_duplicates_include_used_rows_without_secret_diagnostics(self):
        for field in ("uuid", "key"):
            rows = [tracked(1, "used"), credential(2)]
            rows[1][field] = rows[0][field]
            self.make_workbook(rows)
            with self.subTest(field=field), self.assertRaisesRegex(batch.BatchError, "duplicate") as caught:
                self.load()
            self.assertIn("rows 2 and 3", str(caught.exception))
            self.assertNotIn(rows[0][field], str(caught.exception))

    def test_tracking_schema_status_and_lifecycle_validation(self):
        invalid = [dict(status="unknown"), dict(status=4), dict(status=" unused"),
                   dict(status="=A1"), dict(stage="firmware"), dict(device_mac="invalid")]
        for field in ("stage", "device_mac", "product_key", "port", "run_id", "updated_at"):
            row = tracked(1, "in_progress")
            row[field] = None
            invalid.append(row)
        for status, stage, error in (("used", "firmware", None), ("used", "complete", "bad"),
                                     ("fail", "complete", "bad"), ("fail", "firmware", None),
                                     ("in_progress", "complete", None),
                                     ("in_progress", "firmware", "bad")):
            row = tracked(1, status)
            row.update(stage=stage, error=error)
            invalid.append(row)
        for update in invalid:
            with self.subTest(update=update):
                row = credential(1)
                row.update(update)
                self.make_workbook([row])
                with self.assertRaises(batch.BatchError):
                    self.load()

    def test_duplicate_normalized_chip_bindings_rejected(self):
        rows = [tracked(1, "used"), tracked(2, "fail")]
        rows[1]["device_mac"] = "0200.0000.0001"
        self.make_workbook(rows)
        with self.assertRaisesRegex(batch.BatchError, "duplicate device_mac"):
            self.load()

    def test_allocation_sheet_order_never_recycles_consumed_rows(self):
        rows = [tracked(1, "used"), tracked(2, "fail"), tracked(3, "in_progress"),
                credential(4), dict(credential(5), status="unused")]
        self.make_workbook(rows)
        workbook = self.load()
        workbook.rows = dict(reversed(list(workbook.rows.items())))
        pairs = batch.allocate_rows(workbook.rows, list(reversed(self.candidates)))
        self.assertEqual([(device.path, row.index) for device, row in pairs],
                         [(self.candidates[1].path, 5), (self.candidates[0].path, 6)])
        with self.assertRaisesRegex(batch.BatchError, "not enough"):
            batch.allocate_rows(workbook.rows, self.candidates * 2)

    def test_preserves_credentials_other_sheets_columns_and_formatting(self):
        self.make_workbook([dict(credential(1), note="operator note"), credential(2)],
                           ["uuid", "key", "note"])
        workbook = openpyxl.load_workbook(self.xlsx)
        workbook.active["A2"].font = Font(bold=True, color="00112233")
        workbook.active["C2"].fill = PatternFill("solid", fgColor="00ABCDEF")
        workbook.create_sheet("Notes")["B3"] = "=1+2"
        workbook.save(self.xlsx)
        workbook.close()
        assignments, coordinator = self.assignments_and_coordinator()
        coordinator.finish(assignments[0], "used", "complete", None)
        coordinator.finish(assignments[1], "fail", "auth_read", "timeout")
        saved = self.load()
        self.assertEqual(saved.wb["Notes"]["B3"].value, "=1+2")
        self.assertTrue(saved.ws["A2"].font.bold)
        self.assertEqual(saved.ws["C2"].fill.fgColor.rgb, "00ABCDEF")
        self.assertEqual(saved.ws["C2"].value, "operator note")
        for index in (1, 2):
            self.assertEqual(saved.rows[index + 1].uuid, credential(index)["uuid"])
            self.assertEqual(saved.rows[index + 1].key, credential(index)["key"])
        self.assertEqual(saved.ws.cell(2, saved.header["status"]).fill.fgColor.rgb[-6:], "C6EFCE")
        self.assertEqual(saved.ws.cell(3, saved.header["status"]).fill.fgColor.rgb[-6:], "FFC7CE")
        self.assertEqual(saved.ws["B2"].fill.fill_type, None)

    def test_retry_failed_interrupted_rows_original_pid_and_renamed_port(self):
        for status in ("fail", "in_progress"):
            self.make_workbook([tracked(1, status)])
            args = self.args("--retry-uuid", credential(1)["uuid"], "--device", "/dev/renamed")
            workbook = self.load()
            candidate = batch.PortCandidate("/dev/renamed")
            assignment = batch.plan_assignments(args, workbook, [candidate])[0]
            self.assertEqual(assignment.row.index, 2)
            self.assertEqual(assignment.retry_mac, tracked(1, status)["device_mac"])
            batch.validate_probes({candidate.path: batch.ProbeResult("esp32s3", assignment.retry_mac)},
                                  types.SimpleNamespace(chip="esp32s3"), workbook, args)
            with self.assertRaisesRegex(batch.BatchError, "does not match"):
                batch.validate_probes({candidate.path: batch.ProbeResult("esp32s3", "02:00:00:00:00:ff")},
                                      types.SimpleNamespace(chip="esp32s3"), workbook, args)
            args.pid = "different-pid-001"
            with self.assertRaisesRegex(batch.BatchError, "original --pid"):
                batch.plan_assignments(args, workbook, [candidate])

    def test_retry_rejects_used_unused_unknown_uuid_and_wrong_device_count(self):
        for status in ("used", "unused"):
            self.make_workbook([tracked(1, status) if status == "used" else credential(1)])
            with self.subTest(status=status), self.assertRaises(batch.BatchError):
                batch.plan_assignments(self.args("--retry-uuid", credential(1)["uuid"]),
                                       self.load(), self.candidates[:1])
        with self.assertRaises(batch.BatchError):
            batch.plan_assignments(self.args("--retry-uuid", "missing"), self.load(), self.candidates)
        for devices in ([], ["/dev/a", "/dev/b"]):
            args = self.args()
            args.device = devices
            with self.assertRaisesRegex(batch.BatchError, "exactly one"):
                batch._validate_retry_structure(args)

    def test_companion_lock_survives_atomic_replace_and_releases(self):
        first, second = batch.WorkbookLock(str(self.xlsx)), batch.WorkbookLock(str(self.xlsx))
        first.acquire()
        self.addCleanup(first.release)
        self.addCleanup(second.release)
        self.load(write=True).save()
        with self.assertRaisesRegex(batch.BatchError, "another batch"):
            second.acquire()
        first.release()
        second.acquire()
        self.assertEqual(Path(first.path).stat().st_mode & 0o777, 0o600)

    def test_atomic_save_same_directory_fsync_and_permissions(self):
        workbook = self.load(write=True)
        real_replace, real_fsync = os.replace, os.fsync
        events = []

        def replace(source, target):
            self.assertEqual(Path(source).parent, self.xlsx.parent)
            self.assertEqual(Path(source).suffix, ".xlsx")
            self.assertEqual(Path(source).stat().st_mode & 0o777, 0o640)
            self.assertIn("fsync", events)
            events.append("replace")
            return real_replace(source, target)

        def fsync(fd):
            events.append("fsync")
            return real_fsync(fd)

        with mock.patch.object(batch.os, "replace", side_effect=replace), \
             mock.patch.object(batch.os, "fsync", side_effect=fsync):
            workbook.save()
        self.assertEqual(events[-2:], ["replace", "fsync"])
        self.assertEqual(self.xlsx.stat().st_mode & 0o777, 0o640)
        self.assertFalse(list(self.root.glob(".batch_*.xlsx")))

    def test_failed_replace_preserves_original_and_cleans_temporary_file(self):
        workbook = self.load(write=True)
        original = self.xlsx.read_bytes()
        with mock.patch.object(batch.os, "replace", side_effect=OSError("disk failure")), \
             self.assertRaises(OSError):
            workbook.save()
        self.assertEqual(self.xlsx.read_bytes(), original)
        self.assertFalse(list(self.root.glob(".batch_*.xlsx")))

    def test_external_edit_before_or_during_save_is_not_overwritten(self):
        for during in (False, True):
            with self.subTest(during=during):
                self.make_workbook([credential(1)])
                workbook = self.load(write=True)
                original_save = workbook.wb.save
                external = self.xlsx.read_bytes() + b"external edit"

                def save(path):
                    original_save(path)
                    self.xlsx.write_bytes(external)

                if not during:
                    self.xlsx.write_bytes(external)
                with mock.patch.object(workbook.wb, "save", side_effect=save), \
                     self.assertRaisesRegex(batch.PersistenceError, "changed on disk"):
                    workbook.save()
                self.assertEqual(self.xlsx.read_bytes(), external)
                self.assertFalse(list(self.root.glob(".batch_*.xlsx")))

    def test_protected_backup_and_append_only_secret_free_journal(self):
        original = self.xlsx.read_bytes()
        backup = Path(batch.write_backup(str(self.xlsx)))
        self.assertEqual(backup.read_bytes(), original)
        self.assertEqual(backup.stat().st_mode & 0o777, 0o600)
        assignments, coordinator = self.assignments_and_coordinator()
        journal_path = Path(coordinator.journal.path)
        prefix = journal_path.read_bytes()
        coordinator.finish(assignments[0], "fail", "auth_read", assignments[0].row.key + " timeout")
        self.assertTrue(journal_path.read_bytes().startswith(prefix))
        self.assertEqual(journal_path.stat().st_mode & 0o777, 0o600)
        for assignment in assignments:
            self.assertNotIn(assignment.row.key, journal_path.read_text())
        outcome = self.journal_records()[-1]
        self.assertEqual((outcome["event"], outcome["sheet"], outcome["row"], outcome["uuid"]),
                         ("outcome", "Credentials", 2, credential(1)["uuid"]))
        self.assertEqual(self.load().rows[2].error, "[redacted] timeout")


class TestBatchFlow(SyntheticCase):
    def test_dry_run_validates_without_mutation_or_device_access(self):
        original = self.xlsx.read_bytes()
        files = set(self.root.iterdir())
        with mock.patch.object(batch, "probe_device") as probe, \
             mock.patch.object(batch.Workbook, "save") as save, \
             mock.patch.object(batch.WorkbookLock, "acquire") as lock:
            self.assertEqual(batch.main(self.argv("--dry-run")), 0)
        probe.assert_not_called()
        save.assert_not_called()
        lock.assert_not_called()
        self.commands.assert_not_called()
        self.auth.assert_not_called()
        self.assertEqual(self.xlsx.read_bytes(), original)
        self.assertEqual(set(self.root.iterdir()), files)
        self.assertIn("no ports opened, no probing", self.output.getvalue())
        for candidate in self.candidates:
            self.assertIn(candidate.path, self.output.getvalue())
        self.assertNotIn(credential(1)["key"], self.output.getvalue())

    def test_list_ports_needs_no_pid_workbook_build_or_spreadsheet(self):
        with mock.patch.object(batch, "Workbook", side_effect=AssertionError("workbook access")), \
             mock.patch.object(batch, "Manifest", side_effect=AssertionError("build access")), \
             mock.patch.object(batch, "_import_openpyxl", side_effect=AssertionError("dependency")), \
             mock.patch.object(batch, "_validate_idf_tools", side_effect=AssertionError("tool access")):
            self.assertEqual(batch.main(["--list-ports"]), 0)
        self.commands.assert_not_called()
        self.assertIn("1234:5678", self.output.getvalue())
        self.assertIn("Synthetic USB bridge", self.output.getvalue())

    def test_no_devices_is_preflight_failure_without_consumption(self):
        self.discovery.return_value = ([], [])
        original = self.xlsx.read_bytes()
        self.assertEqual(batch.main(self.argv()), 2)
        self.assertEqual(self.xlsx.read_bytes(), original)
        self.assertIn("no USB serial candidates", self.errors.getvalue())
        self.assertIn("--device", self.errors.getvalue())
        self.commands.assert_not_called()

    def test_invalid_local_inputs_never_probe(self):
        with mock.patch.object(batch, "probe_device") as probe:
            for pid in ("", "short", " " + PID, "a" * 32, "a" * 16 + "\x00"):
                with self.subTest(pid=pid):
                    self.assertEqual(batch.main(self.argv("--pid", pid)), 2)
            self.make_workbook([credential(1)])
            self.assertEqual(batch.main(self.argv()), 2)
            self.make_workbook([credential(1), credential(2)])
            (self.build / "bootloader.bin").unlink()
            self.assertEqual(batch.main(self.argv()), 2)
        probe.assert_not_called()
        self.commands.assert_not_called()

    def test_confirmation_precedes_probe_and_noninteractive_requires_yes(self):
        self.enable_devices()
        args = self.argv()
        args.remove("--yes")
        with mock.patch.object(batch.sys.stdin, "isatty", return_value=False):
            self.assertEqual(batch.main(args), 2)
        self.probes.assert_not_called()
        with mock.patch.object(batch.sys, "stdin", io.StringIO("no\n")), \
             mock.patch.object(batch.sys.stdin, "isatty", return_value=True), \
             mock.patch.object(batch.select, "select", return_value=([batch.sys.stdin], [], [])):
            self.assertEqual(batch.main(args), 2)
        self.probes.assert_not_called()
        self.assertIn("replaces the entire NVS", self.output.getvalue())
        self.assertIn("may reset", self.output.getvalue())

    def test_probe_barrier_waits_for_every_candidate_before_durable_reservation(self):
        self.enable_devices()
        both_probing = threading.Barrier(2, timeout=5)
        both_flashing = threading.Barrier(2, timeout=5)
        completed = set()
        lock = threading.Lock()
        original = self.xlsx.read_bytes()

        def probe(candidate, manifest, baud, timeout, event):
            if candidate.path not in completed:
                self.assertEqual(self.xlsx.read_bytes(), original)
                self.assertEqual(self.commands.call_count, 0)
                both_probing.wait()
                with lock:
                    completed.add(candidate.path)
            return self.good_probe(candidate, manifest, baud, timeout, event)

        def flash(command, stage, **kwargs):
            self.assertEqual(completed, {c.path for c in self.candidates})
            saved = self.load()
            self.assertTrue(all(saved.rows[i].status == "in_progress" for i in (2, 3)))
            self.assertEqual([saved.rows[i].device_mac for i in (2, 3)],
                             ["02:00:00:00:00:01", "02:00:00:00:00:02"])
            both_flashing.wait()
            return types.SimpleNamespace(stdout="", stderr="")

        self.probes.side_effect = probe
        self.commands.side_effect = flash
        self.assertEqual(batch.main(self.argv()), 0)
        self.assertEqual(self.commands.call_count, 2)
        self.assertEqual([self.load().rows[i].status for i in (2, 3)], ["used", "used"])
        self.discovery.assert_called_once()

    def test_failed_wrong_duplicate_and_bound_probes_abort_without_writes(self):
        self.enable_devices()
        original_probe = self.good_probe
        for failure in ("busy", "wrong", "duplicate", "bound"):
            with self.subTest(failure=failure):
                rows = [credential(1), credential(2)]
                if failure == "bound":
                    rows.append(tracked(3, "used"))
                self.make_workbook(rows)
                before = self.xlsx.read_bytes()

                def probe(candidate, manifest, baud, timeout, event):
                    result = original_probe(candidate, manifest, baud, timeout, event)
                    if candidate == self.candidates[1]:
                        if failure == "busy":
                            result.error = "port is busy"
                        elif failure == "wrong":
                            result.chip = "ESP32-C3"
                        elif failure == "duplicate":
                            result.mac = "02:00:00:00:00:01"
                        else:
                            result.mac = "02:00:00:00:00:03"
                    return result

                self.probes.side_effect = probe
                self.assertEqual(batch.main(self.argv()), 2)
                self.assertEqual(self.xlsx.read_bytes(), before)
                self.commands.assert_not_called()
                self.auth.assert_not_called()
        self.assertIn(self.candidates[1].path, self.errors.getvalue())

    def test_already_bound_failed_and_interrupted_chips_are_rejected(self):
        for status in ("used", "fail", "in_progress"):
            self.make_workbook([credential(2), tracked(1, status)])
            with self.subTest(status=status), self.assertRaisesRegex(batch.BatchError, "already bound"):
                batch.validate_probes({self.candidates[0].path:
                                       batch.ProbeResult("esp32s3", "02:00:00:00:00:01")},
                                      types.SimpleNamespace(chip="esp32s3"), self.load(), self.args())

    def test_pipeline_order_all_images_timeouts_and_durable_outcomes(self):
        self.enable_devices()
        self.assertEqual(batch.main(self.argv("--jobs", "1", "--timeout", "13", "--baud", "115200")), 0)
        records = self.journal_records()
        for row in (2, 3):
            stages = [r["stage"] for r in records if r["row"] == row]
            self.assertEqual(stages, ["reserved", "firmware", "auth_write", "auth_write",
                                      "auth_read", "auth_verify", "auth_verify", "complete"])
        for call in self.commands.call_args_list:
            self.assertEqual(call.args[1], "firmware")
            self.assertEqual(call.kwargs["timeout"], 13)
            self.assertEqual(len(call.args[0][-10:]), 10)
            for name in self.manifest_data["flash_files"].values():
                self.assertIn(str(self.build / name), call.args[0])
        for call in self.auth.call_args_list:
            self.assertEqual(call.args[1], 115200)
            self.assertEqual(call.kwargs["timeout"], 13)
        self.assertEqual(self.load().rows[2].port, self.candidates[0].path)

    def test_device_stage_failures_stop_pipeline_but_not_healthy_peer(self):
        self.enable_devices()
        for stage in ("firmware", "auth_write", "auth_read", "auth_verify"):
            with self.subTest(stage=stage):
                self.make_workbook([credential(1), credential(2)])
                self.commands.reset_mock()
                self.auth.reset_mock()
                failure = batch.AuthToolError(stage, "synthetic timeout or mismatch", credential(1)["key"])

                def command(argv, current_stage, **kwargs):
                    if stage == "firmware" and self.candidates[0].path in argv:
                        raise failure
                    return types.SimpleNamespace(stdout="", stderr="")

                def auth(path, baud, values, **kwargs):
                    if path == self.candidates[0].path:
                        kwargs["stage_callback"](stage)
                        raise failure
                    return self.good_auth(path, baud, values, **kwargs)

                self.commands.side_effect = command
                self.auth.side_effect = auth
                self.assertEqual(batch.main(self.argv("--jobs", "1")), 1)
                saved = self.load()
                self.assertEqual((saved.rows[2].status, saved.rows[2].stage), ("fail", stage))
                self.assertEqual(saved.rows[3].status, "used")
                self.assertNotIn(credential(1)["key"], saved.rows[2].error)
                self.assertNotIn(credential(1)["key"], self.output.getvalue() + self.errors.getvalue())
                if stage == "firmware":
                    self.assertEqual([call.args[0] for call in self.auth.call_args_list],
                                     [self.candidates[1].path])

    def test_identity_change_before_each_write_blocks_that_stage(self):
        self.enable_devices()
        for change_at in (2, 3):
            with self.subTest(change_at=change_at):
                self.make_workbook([credential(1), credential(2)])
                counts = {}
                self.commands.reset_mock()
                self.auth.reset_mock()

                def probe(candidate, manifest, baud, timeout, event):
                    counts[candidate.path] = counts.get(candidate.path, 0) + 1
                    result = self.good_probe(candidate, manifest, baud, timeout, event)
                    if candidate == self.candidates[0] and counts[candidate.path] == change_at:
                        result.mac = "02:00:00:00:00:ff"
                    return result

                self.probes.side_effect = probe
                self.assertEqual(batch.main(self.argv("--jobs", "1")), 1)
                self.assertEqual(self.commands.call_count, 1 if change_at == 2 else 2)
                self.assertEqual(self.auth.call_count, 1)
                self.assertEqual(self.load().rows[2].stage,
                                 "firmware" if change_at == 2 else "auth_write")

    def test_frozen_snapshot_failure_before_reservation_leaves_rows_unused(self):
        self.enable_devices()
        before = self.xlsx.read_bytes()
        with mock.patch.object(batch, "verify_candidates_frozen",
                               side_effect=batch.BatchError("preflight", "device disappeared")):
            self.assertEqual(batch.main(self.argv()), 2)
        self.assertEqual(self.xlsx.read_bytes(), before)
        self.commands.assert_not_called()

    def test_probe_command_parses_chip_mac_and_forwards_controls(self):
        self.commands.side_effect = None
        self.commands.return_value = types.SimpleNamespace(
            stdout="Chip is ESP32-S3 (revision v0.2)\nMAC: 02:AB:CD:EF:00:01\n", stderr="")
        event = threading.Event()
        result = batch.probe_device(self.candidates[0], batch.Manifest(str(self.build)).load(),
                                    115200, 9, event)
        self.assertEqual((result.chip, result.mac, result.error),
                         ("ESP32-S3", "02:ab:cd:ef:00:01", None))
        command = self.commands.call_args.args[0]
        self.assertEqual(command[-1], "read_mac")
        self.assertIn("--no-stub", command)
        self.assertEqual(command[command.index("--port") + 1], self.candidates[0].path)
        self.assertEqual(command[command.index("--after") + 1], "no_reset")
        self.assertEqual(self.commands.call_args.kwargs, {"timeout": 9, "cancel_event": event})

    def test_probe_transport_failure_malformed_output_and_cancellation(self):
        manifest = batch.Manifest(str(self.build)).load()
        self.commands.side_effect = batch.AuthToolError("probe", "timeout", "port busy")
        self.assertEqual(batch.probe_device(self.candidates[0], manifest, None, 9,
                                           threading.Event()).error, "port busy")
        self.commands.side_effect = None
        self.commands.return_value = types.SimpleNamespace(stdout="not a chip response", stderr="")
        self.assertIn("could not read chip MAC", batch.probe_device(
            self.candidates[0], manifest, None, 9, threading.Event()).error)
        self.commands.side_effect = batch.CommandCancelled("probe")
        with self.assertRaises(batch.CommandCancelled):
            batch.probe_device(self.candidates[0], manifest, None, 9, threading.Event())

    def test_probe_timeout_is_capped_and_chip_compatibility_is_not_a_prefix_match(self):
        self.enable_devices()
        batch.probe_barrier(self.candidates, self.args("--timeout", "999"),
                            batch.Manifest(str(self.build)).load(), threading.Event())
        self.assertEqual([call.args[3] for call in self.probes.call_args_list], [60, 60])
        self.assertTrue(batch.chips_compatible("esp32", "ESP32-D0WDQ6"))
        self.assertTrue(batch.chips_compatible("esp32s3", "ESP32-S3"))
        self.assertFalse(batch.chips_compatible("esp32", "ESP32-S3"))
        self.assertFalse(batch.chips_compatible("esp32s3", None))

    def test_report_full_distinct_paths_metadata_and_plain_text(self):
        assignments, coordinator = self.assignments_and_coordinator()
        results = {}
        for assignment in assignments:
            results[assignment.candidate.path] = batch.WorkResult(
                assignment.candidate.path, assignment.row.index, "FAIL", "auth_read", "timeout")
        batch.print_report(assignments, results)
        text = self.output.getvalue()
        summary = text.split("FAILED DEVICES:")[1]
        for assignment in assignments:
            self.assertIn(assignment.candidate.path, summary)
            self.assertIn(assignment.candidate.serial_number, summary)
            self.assertIn(assignment.candidate.location, summary)
            self.assertIn(assignment.mac, summary)
        self.assertNotIn("\x1b", text)
        self.assertNotIn(credential(1)["key"], text)

    def test_invalid_job_timeout_baud_options_and_global_port(self):
        for option, value in (("--jobs", "0"), ("--jobs", "-1"), ("--timeout", "0"),
                              ("--timeout", "-2"), ("--baud", "0")):
            with self.subTest(option=option, value=value), self.assertRaises(SystemExit) as caught:
                self.args(option, value)
            self.assertEqual(caught.exception.code, 2)
        self.assertEqual(batch.main(self.argv("--port", "/dev/wrong")), 2)
        self.assertIn("ESPPORT", self.errors.getvalue())
        self.commands.assert_not_called()


class TestConcurrencyAndRecovery(SyntheticCase):
    def test_worker_job_bounds_parallelism_and_stable_mapping(self):
        self.enable_devices()
        for requested in (1, 2, 99):
            with self.subTest(jobs=requested):
                self.make_workbook([credential(1), credential(2)])
                assignments, coordinator = self.assignments_and_coordinator()
                width = min(requested, len(assignments))
                rendezvous = threading.Barrier(width, timeout=5)
                lock = threading.Lock()
                active = set()
                maximum = [0]

                def command(argv, stage, **kwargs):
                    path = argv[argv.index("--port") + 1]
                    with lock:
                        self.assertNotIn(path, active)
                        active.add(path)
                        maximum[0] = max(maximum[0], len(active))
                    try:
                        rendezvous.wait()
                    finally:
                        with lock:
                            active.remove(path)
                    return types.SimpleNamespace(stdout="", stderr="")

                self.commands.side_effect = command
                results = batch.run_workers(assignments, batch.Manifest(str(self.build)).load(),
                                            coordinator, requested, None, 7)
                self.assertEqual(maximum[0], width)
                self.assertTrue(all(result.result == "PASS" for result in results.values()))
                saved = self.load()
                self.assertEqual([(saved.rows[i].uuid, saved.rows[i].port) for i in (2, 3)],
                                 [(credential(i)["uuid"], self.candidates[i - 1].path) for i in (1, 2)])

    def test_reverse_completion_order_preserves_assignments_and_serializes_saves(self):
        self.enable_devices()
        assignments, coordinator = self.assignments_and_coordinator()
        second_saved = threading.Event()
        save_lock = threading.Lock()
        original_save = coordinator.workbook.save
        original_finish = coordinator.finish
        order = []

        def auth(path, baud, values, **kwargs):
            if path == self.candidates[0].path:
                self.assertTrue(second_saved.wait(5), "second device never completed")
            return self.good_auth(path, baud, values, **kwargs)

        def save():
            self.assertTrue(save_lock.acquire(blocking=False), "concurrent workbook save")
            try:
                original_save()
            finally:
                save_lock.release()

        def finish(assignment, *args):
            original_finish(assignment, *args)
            order.append(assignment.row.index)
            if assignment.row.index == 3:
                second_saved.set()

        self.auth.side_effect = auth
        with mock.patch.object(coordinator.workbook, "save", side_effect=save), \
             mock.patch.object(coordinator, "finish", side_effect=finish):
            results = batch.run_workers(assignments, batch.Manifest(str(self.build)).load(),
                                        coordinator, 2, None, 7)
        self.assertEqual(order, [3, 2])
        self.assertEqual(results[self.candidates[0].path].row, 2)
        self.assertEqual(results[self.candidates[1].path].row, 3)
        self.assertEqual(self.load().rows[2].device_mac, assignments[0].mac)

    def test_journal_is_flushed_before_corresponding_workbook_save(self):
        assignments, coordinator = self.assignments_and_coordinator()
        original_save = coordinator.workbook.save

        def save():
            outcome = self.journal_records()[-1]
            self.assertEqual((outcome["event"], outcome["status"]), ("outcome", "used"))
            return original_save()

        with mock.patch.object(coordinator.workbook, "save", side_effect=save):
            coordinator.finish(assignments[0], "used", "complete", None)

    def test_final_save_failure_stops_queued_peers_retains_durable_reservations(self):
        self.enable_devices()
        assignments, coordinator = self.assignments_and_coordinator()
        with mock.patch.object(coordinator.workbook, "save", side_effect=OSError("disk full")):
            results = batch.run_workers(assignments, batch.Manifest(str(self.build)).load(),
                                        coordinator, 1, None, 7)
        self.assertEqual(results[self.candidates[0].path].result, "RESULT NOT SAVED")
        self.assertEqual(results[self.candidates[1].path].result, "CANCELLED")
        self.assertTrue(coordinator.cancel_event.is_set())
        self.assertIsNotNone(coordinator.failure)
        self.assertEqual(self.commands.call_count, 1)
        self.assertEqual([self.load().rows[i].status for i in (2, 3)], ["in_progress", "in_progress"])
        self.assertEqual(self.journal_records()[-1]["status"], "used")

    def test_journal_failure_stops_operations_without_claiming_saved_failure(self):
        self.enable_devices()
        assignments, coordinator = self.assignments_and_coordinator()
        with mock.patch.object(coordinator.journal, "append", side_effect=OSError("journal full")):
            result = batch.flash_device(assignments[0], batch.Manifest(str(self.build)).load(),
                                        coordinator, None, 7)
        self.assertEqual(result.result, "RESULT NOT SAVED")
        self.assertTrue(coordinator.cancel_event.is_set())
        self.assertEqual(self.load().rows[2].status, "in_progress")
        self.commands.assert_not_called()

    def test_reservation_save_failure_never_launches_workers(self):
        self.enable_devices()
        before = self.xlsx.read_bytes()
        with mock.patch.object(batch.Workbook, "save", side_effect=OSError("disk full")):
            self.assertEqual(batch.main(self.argv()), 1)
        self.assertEqual(self.xlsx.read_bytes(), before)
        self.assertIn("RESULT NOT SAVED", self.errors.getvalue())
        self.commands.assert_not_called()
        self.auth.assert_not_called()
        self.assertEqual([record["event"] for record in self.journal_records()], ["reserve", "reserve"])

    def test_final_save_failure_cli_returns_one_redacts_and_releases_lock(self):
        self.enable_devices()
        original_save = batch.Workbook.save
        saves = []

        def save(workbook):
            saves.append(None)
            if len(saves) > 1:
                raise OSError("save failed " + credential(1)["key"])
            return original_save(workbook)

        with mock.patch.object(batch.Workbook, "save", autospec=True, side_effect=save):
            self.assertEqual(batch.main(self.argv("--jobs", "1")), 1)
        self.assertIn("RESULT NOT SAVED", self.errors.getvalue())
        self.assertNotIn(credential(1)["key"], self.errors.getvalue() + self.output.getvalue())
        self.assertEqual([self.load().rows[i].status for i in (2, 3)], ["in_progress", "in_progress"])
        lock = batch.WorkbookLock(str(self.xlsx))
        lock.acquire()
        lock.release()
        self.assertEqual(self.commands.call_count, 1)

    def test_lock_held_during_confirmation_and_probe(self):
        self.enable_devices()
        competing = batch.WorkbookLock(str(self.xlsx))
        self.addCleanup(competing.release)
        checks = []

        def check_lock(*_args):
            contender = batch.WorkbookLock(str(self.xlsx))
            try:
                with self.assertRaisesRegex(batch.BatchError, "another batch"):
                    contender.acquire()
                checks.append(True)
            finally:
                contender.release()

        def probe(*args):
            check_lock()
            return self.good_probe(*args)

        self.probes.side_effect = probe
        args = self.argv()
        args.remove("--yes")
        with mock.patch.object(batch, "_confirm", side_effect=check_lock):
            self.assertEqual(batch.main(args), 0)
        self.assertEqual(len(checks), 7)
        competing.acquire()
        competing.release()

    def test_persistence_failure_cancels_active_peer(self):
        self.enable_devices()
        assignments, coordinator = self.assignments_and_coordinator()
        active = threading.Barrier(2, timeout=5)

        def command(argv, stage, **kwargs):
            active.wait()
            if self.candidates[1].path in argv:
                self.assertTrue(kwargs["cancel_event"].wait(5), "peer did not cancel")
                raise batch.CommandCancelled(stage)
            return types.SimpleNamespace(stdout="", stderr="")

        self.commands.side_effect = command
        with mock.patch.object(coordinator.workbook, "save", side_effect=OSError("disk full")):
            results = batch.run_workers(assignments, batch.Manifest(str(self.build)).load(),
                                        coordinator, 2, None, 7)
        self.assertEqual(results[self.candidates[0].path].result, "RESULT NOT SAVED")
        self.assertEqual(results[self.candidates[1].path].result, "CANCELLED")
        self.assertEqual(self.auth.call_count, 1)

    def test_cancel_during_firmware_returns_130_and_preserves_nonreusable_rows(self):
        self.enable_devices()

        def command(argv, stage, **kwargs):
            kwargs["cancel_event"].set()
            raise batch.CommandCancelled(stage)

        self.commands.side_effect = command
        self.assertEqual(batch.main(self.argv("--jobs", "1")), 130)
        self.assertEqual(self.commands.call_count, 1)
        self.auth.assert_not_called()
        saved = self.load()
        self.assertEqual([saved.rows[i].status for i in (2, 3)], ["in_progress", "in_progress"])
        pairs = batch.allocate_rows(saved.rows, self.candidates)
        self.assertEqual([row.index for _, row in pairs], [4, 5])

    def test_cancel_during_probe_consumes_nothing(self):
        self.enable_devices()
        before = self.xlsx.read_bytes()

        def probe(candidate, manifest, baud, timeout, event):
            event.set()
            raise batch.CommandCancelled("probe")

        self.probes.side_effect = probe
        self.assertEqual(batch.main(self.argv()), 130)
        self.assertEqual(self.xlsx.read_bytes(), before)
        self.commands.assert_not_called()

    def test_keyboard_interrupt_and_signal_handlers_set_cancellation(self):
        with mock.patch.object(batch, "cmd_flash", side_effect=KeyboardInterrupt):
            self.assertEqual(batch.main(self.argv()), 130)
        event = threading.Event()
        with mock.patch.object(batch.signal, "signal") as signal:
            ORIGINAL_INSTALL_SIGNALS(event)
        self.assertEqual(signal.call_count, 2)
        for call in signal.call_args_list:
            event.clear()
            call.args[1](call.args[0], None)
            call.args[1](call.args[0], None)
            self.assertTrue(event.is_set())


class TestProgressDisplay(SyntheticCase):
    def test_plain_run_keeps_prefixes_has_no_ansi_and_summarizes(self):
        self.enable_devices()
        self.assertEqual(batch.main(self.argv()), 0)
        output = self.output.getvalue()
        for candidate in self.candidates:
            self.assertIn("[auth_read] " + candidate.path, output)
        self.assertNotIn("\x1b", output)
        self.assertIn("Probing complete: responded 2/2 | errors 0", output)
        self.assertIn("Flashing complete: finished 2/2 | pass 2 | fail 0 | "
                      "cancelled 0 | unsaved 0", output)
        self.assertNotIn(credential(1)["key"], output + self.errors.getvalue())

    def test_live_frames_render_counts_labels_and_dedupe(self):
        paths = ["/dev/cu.usbserial-%03d" % i for i in (1, 2)]
        stream = FakeTTY()
        progress = batch.ProgressDisplay(display_assignments(paths), stream=stream,
                                         size=(100, 30), live=True)
        now = [100.0]
        progress._clock = lambda: now[0]
        progress.start("probing")
        now[0] += 1
        for path in paths:
            progress.probe_start(path)
        progress.tick(force=True)
        frame = stream.getvalue()
        self.assertIn("probing | responded 0/2 | probing 2 | waiting 0 | errors 0", frame)
        for path in paths:
            self.assertIn(path, frame)
        now[0] += 2
        progress.probe(paths[0], True)
        progress.probe(paths[1], False)
        progress.tick(force=True)
        frame = stream.getvalue()
        self.assertIn("responded 2/2", frame)
        self.assertIn("errors 1", frame)
        self.assertIn("probe error", frame)
        progress.stop()
        progress.start("flashing")
        now[0] += 1
        progress.stage(paths[0], "firmware")
        progress.tick(force=True)
        progress.stage(paths[0], "firmware")  # duplicate must not reset timing
        now[0] += 5
        progress.tick(force=True)
        self.assertIn("00:05", stream.getvalue())
        progress.finish(paths[0], 2, batch.RESULT_PASS)
        progress.tick(force=True)
        frame = stream.getvalue()
        self.assertIn("pass 1", frame)
        self.assertIn("PASS (verified)", frame)
        progress.stop()
        self.assertTrue(stream.getvalue().endswith("\n"))

    def test_live_frames_never_reach_the_full_terminal_width(self):
        paths = ["/dev/cu.usbserial-long-name-%03d" % i for i in range(1, 4)]
        for width in (60, 80):
            with self.subTest(width=width):
                progress = batch.ProgressDisplay(display_assignments(paths),
                                                 stream=FakeTTY(), size=(width, 30),
                                                 live=True)
                progress._clock = lambda: 100.0
                progress.start("flashing")
                for path in paths:
                    progress.stage(path, "auth_read")
                lines = progress._build_lines(100.0)
                self.assertLessEqual(max(len(line) for line in lines), width - 1)

    def test_live_falls_back_for_small_terminal_dumb_or_overflow(self):
        paths = ["/dev/cu.p%03d" % i for i in range(12)]
        for size in ((100, 8), (40, 30)):
            with self.subTest(size=size):
                progress = batch.ProgressDisplay(display_assignments(paths),
                                                 stream=FakeTTY(), size=size, live=True)
                self.assertFalse(progress._live)
        with mock.patch.dict(os.environ, {"TERM": "dumb"}):
            progress = batch.ProgressDisplay(display_assignments(paths[:2]),
                                             stream=FakeTTY(), size=(100, 30))
            self.assertFalse(progress._live)
        progress = batch.ProgressDisplay(display_assignments(paths[:2]),
                                         stream=io.StringIO())
        self.assertFalse(progress._live)

    def test_live_stop_is_idempotent_and_sanitizes_control_characters(self):
        path = "/dev/cu.bad\nname\x1b[2J"
        stream = FakeTTY()
        progress = batch.ProgressDisplay(display_assignments([path]), stream=stream,
                                         size=(100, 30), live=True)
        progress._clock = lambda: 100.0
        progress.start("probing")
        progress.tick(force=True)
        progress.stop()
        progress.stop()
        progress.stage(path, "firmware")
        progress.finish(path, 2, batch.RESULT_PASS)
        rendered = stream.getvalue()
        self.assertNotIn("\nname", rendered)
        self.assertNotIn("\x1b[2J", rendered)
        self.assertIn("?", rendered)

    def test_attached_display_reports_terminal_outcomes_and_closes_cleanly(self):
        self.enable_devices()
        assignments, coordinator = self.assignments_and_coordinator()
        stream = io.StringIO()
        progress = batch.ProgressDisplay(assignments, stream=stream)
        coordinator.progress = progress
        progress.start("flashing")
        results = batch.run_workers(assignments, batch.Manifest(str(self.build)).load(),
                                    coordinator, 2, None, 7)
        progress.stop()
        output = stream.getvalue()
        self.assertTrue(all(result.result == "PASS" for result in results.values()))
        self.assertIn("[PASS] " + self.candidates[0].path, output)
        self.assertIn("Flashing complete: finished 2/2 | pass 2", output)

    def test_attached_display_counts_unsaved_and_cancelled(self):
        self.enable_devices()
        assignments, coordinator = self.assignments_and_coordinator()
        stream = io.StringIO()
        progress = batch.ProgressDisplay(assignments, stream=stream)
        coordinator.progress = progress
        progress.start("flashing")
        with mock.patch.object(coordinator.workbook, "save", side_effect=OSError("disk full")):
            results = batch.run_workers(assignments, batch.Manifest(str(self.build)).load(),
                                        coordinator, 1, None, 7)
        progress.stop()
        output = stream.getvalue()
        self.assertEqual(results[self.candidates[0].path].result, "RESULT NOT SAVED")
        self.assertEqual(results[self.candidates[1].path].result, "CANCELLED")
        self.assertIn("[RESULT NOT SAVED]", output)
        self.assertIn("[CANCELLED]", output)
        self.assertIn("Flashing complete: finished 2/2", output)
        self.assertIn("cancelled 1", output)
        self.assertIn("unsaved 1", output)
        self.assertIn("Stopping; waiting for active commands to exit.", output)

    def test_concurrent_updates_remain_consistent(self):
        paths = ["/dev/cu.c%02d" % i for i in range(4)]
        stream = io.StringIO()
        progress = batch.ProgressDisplay(display_assignments(paths), stream=stream)
        progress.start("flashing")
        barrier = threading.Barrier(len(paths), timeout=5)

        def worker(index):
            progress.stage(paths[index], "firmware")
            barrier.wait()
            progress.finish(paths[index], index + 2, batch.RESULT_PASS)

        threads = [threading.Thread(target=worker, args=(index,))
                   for index in range(len(paths))]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join(timeout=5)
        self.assertTrue(all(not thread.is_alive() for thread in threads))
        progress.stop()
        self.assertIn("Flashing complete: finished 4/4 | pass 4", stream.getvalue())

    def test_rendering_failures_are_contained(self):
        paths = ["/dev/cu.usbserial-001", "/dev/cu.usbserial-002"]
        live = batch.ProgressDisplay(display_assignments(paths), stream=BrokenStream(),
                                     size=(100, 30), live=True)
        live.start("probing")
        live.probe_start(paths[0])
        live.tick(force=True)
        self.assertFalse(live._live)
        live.stop()
        plain = batch.ProgressDisplay(display_assignments(paths), stream=BrokenStream(),
                                      live=False)
        plain.start("probing")
        plain.probe_start(paths[0])
        plain.finish(paths[0], 2, batch.RESULT_PASS)
        plain.stop()

    @staticmethod
    def _last_frame_width(stream):
        last = stream.getvalue().rsplit("\r", 1)[-1]
        return max(len(line.replace("\x1b[K", "")) for line in last.split("\n"))

    def test_journal_close_failure_suppresses_success_footer(self):
        self.enable_devices()
        with mock.patch.object(batch.Journal, "close",
                               side_effect=OSError("journal disk full")):
            self.assertEqual(batch.main(self.argv()), 1)
        output = self.output.getvalue()
        self.assertNotIn("Flashing complete", output)
        self.assertNotIn("All 2 device(s) verified", output)
        self.assertIn("RESULT NOT SAVED", self.errors.getvalue())

    def test_live_resize_rebuilds_frame_at_the_new_width(self):
        paths = ["/dev/cu.usbserial-long-name-%03d" % i for i in (1, 2, 3)]
        stream = FakeTTY()
        progress = batch.ProgressDisplay(display_assignments(paths), stream=stream,
                                         live=True)
        progress._clock = lambda: 100.0
        with mock.patch.object(batch, "_terminal_size", return_value=(100, 30)):
            progress.start("flashing")
            for path in paths:
                progress.stage(path, "auth_read")
            progress.tick(force=True)
            self.assertEqual(progress._width, 100)
            self.assertLessEqual(self._last_frame_width(stream), 99)
            stream.seek(0)
            stream.truncate(0)
            with mock.patch.object(batch, "_terminal_size", return_value=(70, 30)):
                progress.tick(force=True)
            self.assertEqual(progress._width, 70)
            self.assertLessEqual(self._last_frame_width(stream), 69)
            with mock.patch.object(batch, "_terminal_size", return_value=(50, 30)):
                progress.tick(force=True)
        self.assertFalse(progress._live)
        progress.stop()

    def test_live_stop_after_shrink_leaves_clean_plain_output(self):
        paths = ["/dev/cu.usbserial-001", "/dev/cu.usbserial-002"]
        stream = FakeTTY()
        progress = batch.ProgressDisplay(display_assignments(paths), stream=stream,
                                         live=True)
        progress._clock = lambda: 100.0
        with mock.patch.object(batch, "_terminal_size", return_value=(100, 30)):
            progress.start("flashing")
            progress.stage(paths[0], "auth_read")
            progress.tick(force=True)
            with mock.patch.object(batch, "_terminal_size", return_value=(50, 30)):
                progress.stop()
        self.assertFalse(progress._live)
        self.assertIn("Flashing complete", stream.getvalue())

    def test_worker_events_are_queued_until_the_collector_flushes(self):
        paths = ["/dev/cu.usbserial-001", "/dev/cu.usbserial-002"]
        stream = io.StringIO()
        progress = batch.ProgressDisplay(display_assignments(paths), stream=stream)
        progress.start("flashing")
        for path in paths:
            progress.stage(path, "firmware")
        progress.finish(paths[0], 2, batch.RESULT_PASS)
        self.assertEqual(stream.getvalue(), "")
        self.assertEqual(len(progress._pending), 3)
        progress.tick()
        output = stream.getvalue()
        self.assertIn("[firmware] " + paths[0], output)
        self.assertIn("[PASS] " + paths[0], output)
        self.assertEqual(progress._pending, [])

    def test_live_fallback_flushes_subsequent_events_as_plain_lines(self):
        paths = ["/dev/cu.usbserial-001", "/dev/cu.usbserial-002"]
        stream = FakeTTY()
        with mock.patch.object(batch, "_terminal_size", return_value=(100, 30)):
            progress = batch.ProgressDisplay(display_assignments(paths), stream=stream,
                                             live=True)
            progress._clock = lambda: 100.0
            progress.start("probing")
            progress.tick(force=True)
            progress.stage(paths[0], "firmware")
            self.assertEqual(progress._pending, [])
            with mock.patch.object(batch, "_terminal_size", return_value=(40, 30)):
                progress.tick(force=True)
        self.assertFalse(progress._live)
        progress.stage(paths[1], "firmware")
        progress.tick()
        self.assertIn("[firmware] " + paths[1], stream.getvalue())

    def test_probe_barrier_reports_stopping_when_cancelled(self):
        self.enable_devices()
        stream = io.StringIO()
        progress = batch.ProgressDisplay(
            display_assignments([candidate.path for candidate in self.candidates]),
            stream=stream)
        cancel_event = threading.Event()
        cancel_event.set()
        results = batch.probe_barrier(self.candidates, self.args("--timeout", "5"),
                                      batch.Manifest(str(self.build)).load(),
                                      cancel_event, progress)
        self.assertEqual(len(results), len(self.candidates))
        self.assertIn("Stopping; waiting for active commands to exit.",
                      stream.getvalue())

    def test_live_display_closes_before_final_report(self):
        self.enable_devices()
        real = batch.ProgressDisplay

        def factory(assignments, **kwargs):
            return real(assignments, stream=self.output, live=True, size=(100, 30))

        with mock.patch.object(batch, "ProgressDisplay", side_effect=factory):
            self.assertEqual(batch.main(self.argv()), 0)
        output = self.output.getvalue()
        self.assertIn("\x1b[K", output)
        self.assertIn("DEVICE", output)
        self.assertIn("All 2 device(s) verified", output)
        for candidate in self.candidates:
            self.assertIn(candidate.path, output)
        self.assertNotIn(credential(1)["key"], output + self.errors.getvalue())


ORIGINAL_DISCOVER = batch.discover_candidates
ORIGINAL_INSTALL_SIGNALS = batch._install_signal_handlers

if __name__ == "__main__":
    unittest.main()
