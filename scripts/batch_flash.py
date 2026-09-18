#!/usr/bin/env python3
"""Batch firmware + Tuya credential flashing backend (standalone).

Discovers connected USB serial candidates, previews a deterministic
port-to-credential-row mapping, probes every selected device as a preflight
barrier, then flashes the full build manifest image set and writes/inspects the
unique Tuya identity on each device concurrently. Results are persisted to the
credential workbook (and a secret-free journal) as they complete.

The script is usable on its own with the ESP-IDF Python interpreter:

    python scripts/batch_flash.py --list-ports
    python scripts/batch_flash.py --pid YOUR_PRODUCT_PID --dry-run
    python scripts/batch_flash.py --pid YOUR_PRODUCT_PID
    python scripts/batch_flash.py --pid YOUR_PRODUCT_PID --device /dev/cu.usbserial-001

The project idf.py extension forwards global -B/--build-dir and -b/--baud to
this script; a non-empty global -p/--port is rejected (use --device instead).

Exit codes: 0 all jobs verified and saved, 1 device or persistence failure,
2 invalid input/preflight rejection, 130 cancellation.
"""

import argparse
import concurrent.futures
import dataclasses
import errno
import hashlib
import json
import math
import os
import re
import shutil
import select
import signal
import sys
import tempfile
import threading
import time
import uuid
from datetime import datetime, timezone

try:
    import fcntl
except ImportError:  # pragma: no cover - non-POSIX hosts are unsupported
    fcntl = None

_HERE = os.path.dirname(os.path.abspath(__file__))
if _HERE not in sys.path:
    sys.path.insert(0, _HERE)

from tuya_auth_tool import (  # noqa: E402
    AuthToolError,
    CommandCancelled,
    LENGTH_LIMITS,
    check_cancelled,
    find_idf_tools_structured,
    run_command,
    write_verify_identity,
)

REQUIREMENTS = os.path.join(_HERE, "requirements-batch-flash.txt")

CREDENTIAL_HEADERS = ("uuid", "key")

STATUS_UNUSED = "unused"
STATUS_IN_PROGRESS = "in_progress"
STATUS_USED = "used"
STATUS_FAIL = "fail"
VALID_STATUSES = (STATUS_UNUSED, STATUS_IN_PROGRESS, STATUS_USED, STATUS_FAIL)

STAGE_RESERVED = "reserved"
STAGE_FIRMWARE = "firmware"
STAGE_AUTH_WRITE = "auth_write"
STAGE_AUTH_READ = "auth_read"
STAGE_AUTH_VERIFY = "auth_verify"
STAGE_COMPLETE = "complete"
VALID_STAGES = (STAGE_RESERVED, STAGE_FIRMWARE, STAGE_AUTH_WRITE,
                STAGE_AUTH_READ, STAGE_AUTH_VERIFY, STAGE_COMPLETE)

# Tracking columns appended after the credential columns when missing.
TRACKING_FIELDS = ("status", "stage", "error", "port", "usb_serial",
                   "usb_location", "device_mac", "product_key", "run_id",
                   "updated_at")

DEFAULT_TIMEOUT = 600.0
PROBE_TIMEOUT_CAP = 60.0
ERROR_MAX_LEN = 200

RESULT_PASS = "PASS"
RESULT_FAIL = "FAIL"
RESULT_NOT_SAVED = "RESULT NOT SAVED"
RESULT_CANCELLED = "CANCELLED"

_FILL_GREEN = "C6EFCE"
_FILL_RED = "FFC7CE"
_FILL_NEUTRAL = None

_CHIP_RE = re.compile(r"(?:Chip is|Chip type:|Detecting chip type\.\.\.)\s+(ESP\S+)", re.I)
_MAC_RE = re.compile(r"^\s*MAC:\s*((?:[0-9a-fA-F]{2}:){5}[0-9a-fA-F]{2})\s*$", re.M)
_BASE_MAC_RE = re.compile(
    r"^\s*BASE MAC:\s*((?:[0-9a-fA-F]{2}:){5}[0-9a-fA-F]{2})\s*$", re.M)
_MACOS_ALIAS_RE = re.compile(r"^/dev/(?:cu|tty)\.(.+)$")


class BatchError(Exception):
    """User-facing error with an exit code and a stage for reporting."""

    def __init__(self, stage, message, detail=None, exit_code=2):
        super().__init__(message)
        self.stage = stage
        self.message = message
        self.detail = detail
        self.exit_code = exit_code

    def __str__(self):
        if self.detail:
            return "%s: %s" % (self.message, self.detail)
        return self.message


class PersistenceError(BatchError):
    def __init__(self, message, detail=None):
        super().__init__("persistence", message, detail, exit_code=1)


# ---------------------------------------------------------------------------
# USB candidate discovery
# ---------------------------------------------------------------------------

@dataclasses.dataclass
class PortCandidate:
    path: str
    description: str = ""
    hwid: str = ""
    vid: object = None
    pid: object = None
    serial_number: object = None
    location: object = None

    @property
    def vid_pid(self):
        if self.vid is None or self.pid is None:
            return ""
        return "%04x:%04x" % (self.vid, self.pid)


def _import_list_ports():
    try:
        from serial.tools import list_ports
    except ImportError as exc:
        raise BatchError(
            "preflight", "pyserial is not installed",
            "run idf tuya-batch-setup, or install with: %s -m pip install -r %s"
            % (sys.executable, REQUIREMENTS)) from exc
    return list_ports


def _try_list_ports():
    try:
        return _import_list_ports()
    except BatchError:
        return None


def _comports(list_ports):
    try:
        return list(list_ports.comports(include_links=True))
    except TypeError:
        return list(list_ports.comports())


def _alias_key_for_path(path):
    match = _MACOS_ALIAS_RE.match(path)
    if match:
        return "mac:" + match.group(1)
    return "real:" + os.path.realpath(path)


def _alias_preference(candidate):
    # Prefer the macOS callout node (/dev/cu.*) over its call-in alias.
    return (0 if "/dev/cu." in candidate.path else 1, candidate.path)


def _candidate_from_port(port, path=None):
    return PortCandidate(
        path=path or port.device,
        description=getattr(port, "description", "") or "",
        hwid=getattr(port, "hwid", "") or "",
        vid=getattr(port, "vid", None),
        pid=getattr(port, "pid", None),
        serial_number=getattr(port, "serial_number", None),
        location=getattr(port, "location", None),
    )


# macOS callout nodes that are never USB-backed ESP candidates.
_NON_USB_MACOS_MARKERS = ("bluetooth", "debug-console", "wireless", "iap")
# Linux tty names backed by USB serial bridges versus built-in UARTs.
_LINUX_USB_TTY_PREFIXES = ("ttyusb", "ttyacm")
_LINUX_BUILTIN_TTY_PREFIXES = ("ttys", "ttyama", "ttygs", "ttymxc", "ttyps",
                               "ttyprintk", "tty0")

_ENUM_LOCK = threading.Lock()


def _platform_include(port):
    """Return (include, reason) for a listed serial port.

    Only interfaces that are actually USB-backed qualify: arbitrary ports that
    merely expose VID metadata do not, and Bluetooth/built-in UARTs are
    excluded. Linux discovery also accepts ``/dev/serial/by-id|by-path``
    symlinks that resolve to a USB tty.
    """
    device = port.device or ""
    base = os.path.basename(device).lower()
    resolved = os.path.realpath(device) if os.path.islink(device) else device
    resolved_base = os.path.basename(resolved).lower()
    if "bluetooth" in base or "debug-console" in base:
        return False, "bluetooth/debug control port"
    if sys.platform == "darwin":
        if base.startswith("tty."):
            return False, "duplicate call-in alias (/dev/tty.*)"
        if not base.startswith("cu."):
            return False, "not a macOS callout port"
        if any(marker in base for marker in _NON_USB_MACOS_MARKERS):
            return False, "non-USB macOS callout port"
        usb_cues = " ".join(str(value or "") for value in (
            device, getattr(port, "hwid", ""), getattr(port, "location", ""))).lower()
        if getattr(port, "vid", None) is None and "usb" not in usb_cues:
            return False, "no USB metadata"
        return True, ""
    if base.startswith("rfcomm"):
        return False, "bluetooth serial endpoint"
    if base.startswith(_LINUX_BUILTIN_TTY_PREFIXES):
        return False, "built-in UART endpoint"
    if base.startswith(_LINUX_USB_TTY_PREFIXES):
        return True, ""
    if "/dev/serial/by-id/" in device or "/dev/serial/by-path/" in device:
        if resolved_base.startswith(_LINUX_USB_TTY_PREFIXES):
            return True, ""
        return False, "serial symlink is not a USB serial interface"
    return False, "not a USB-backed serial interface"


def _dedupe_candidates(candidates):
    groups = {}
    for candidate in candidates:
        groups.setdefault(_alias_key_for_path(candidate.path), []).append(candidate)
    kept = []
    dropped = []
    for group in groups.values():
        if len(group) == 1:
            kept.append(group[0])
            continue
        chosen = min(group, key=_alias_preference)
        kept.append(chosen)
        for candidate in group:
            if candidate is not chosen:
                dropped.append((candidate.path, "duplicate alias of %s" % chosen.path))
    kept.sort(key=lambda candidate: candidate.path)
    return kept, dropped


def discover_candidates():
    """Enumerate, filter and deduplicate USB serial candidates."""
    list_ports = _import_list_ports()
    included = []
    excluded = []
    try:
        ports = _comports(list_ports)
    except Exception as exc:  # pragma: no cover - platform backend failure
        raise BatchError("discovery", "could not enumerate serial ports", str(exc)) from exc
    for port in ports:
        ok, reason = _platform_include(port)
        if ok:
            included.append(_candidate_from_port(port))
        else:
            excluded.append((port.device, reason))
    deduped, dropped = _dedupe_candidates(included)
    excluded.extend(dropped)
    return deduped, excluded


def _metadata_map():
    list_ports = _try_list_ports()
    if list_ports is None:
        return {}
    result = {}
    for port in _comports(list_ports):
        result[port.device] = port
        result[_alias_key_for_path(port.device)] = port
    return result


def explicit_candidates(paths):
    """Build candidates from explicit --device paths, in argument order."""
    if len(set(paths)) != len(paths):
        raise BatchError("preflight", "duplicate --device path", ", ".join(paths))
    seen = {}
    for path in paths:
        if not os.path.isabs(path):
            raise BatchError("preflight", "--device requires an absolute path", path)
        key = _alias_key_for_path(path)
        if key in seen:
            raise BatchError("preflight", "duplicate --device alias",
                             "%s and %s refer to the same port" % (seen[key], path))
        seen[key] = path
    metadata = _metadata_map()
    candidates = []
    for path in paths:
        if not os.path.exists(path):
            raise BatchError("preflight", "device path not found", path)
        info = metadata.get(path) or metadata.get(_alias_key_for_path(path))
        if info is not None:
            candidates.append(_candidate_from_port(info, path=path))
        else:
            candidates.append(PortCandidate(path=path))
    return candidates


def verify_candidates_frozen(candidates):
    """Re-enumerate selected ports; reject disappearance or identity changes.

    Called after the snapshot is taken (post-preview/probe) and again
    immediately before each firmware write. Paths are never replaced; a
    metadata-less explicit candidate only has to still exist.
    """
    with _ENUM_LOCK:
        metadata = _metadata_map()
        for candidate in candidates:
            if not os.path.exists(candidate.path):
                raise BatchError("preflight", "device path disappeared",
                                 candidate.path)
            info = metadata.get(candidate.path) or metadata.get(
                _alias_key_for_path(candidate.path))
            has_metadata = any(getattr(candidate, field) is not None
                               for field in ("vid", "pid", "serial_number", "location"))
            if info is None:
                if has_metadata:
                    raise BatchError(
                        "preflight", "device metadata disappeared since discovery",
                        candidate.path)
                # Metadata-less explicit candidates can only be checked by path.
                continue
            if not has_metadata:
                continue
            fresh = _candidate_from_port(info, path=candidate.path)
            for field in ("vid", "pid", "serial_number", "location"):
                if getattr(candidate, field) != getattr(fresh, field):
                    raise BatchError(
                        "preflight", "device metadata changed since discovery",
                        "%s (%s changed)" % (candidate.path, field))


def _resolve_candidates(args):
    if args.device:
        return explicit_candidates(args.device)
    candidates, _excluded = discover_candidates()
    if not candidates:
        raise BatchError(
            "discovery", "no USB serial candidates found",
            "check cables/drivers, or pass --device PORT for a path whose "
            "USB metadata is unavailable")
    return candidates


# ---------------------------------------------------------------------------
# Firmware manifest
# ---------------------------------------------------------------------------

_ESTOOL_SUBCOMMANDS = {}
_ESTOOL_SUBCOMMAND_LOCK = threading.Lock()


def esptool_subcommand(name):
    """Return the subcommand spelling accepted by the installed esptool.

    esptool <=4 registers ``write_flash``/``read_mac``; esptool 5 prefers
    ``write-flash``/``read-mac`` (still accepting the underscore aliases but
    warning about them). Probe the CLI once so either dialect works.
    """
    with _ESTOOL_SUBCOMMAND_LOCK:
        cached = _ESTOOL_SUBCOMMANDS.get(name)
        if cached is not None:
            return cached
        chosen = name
        hyphen = name.replace("_", "-")
        try:
            result = run_command([sys.executable, "-m", "esptool", "--help"],
                                 "preflight", timeout=30.0)
            text = (result.stdout or "") + (result.stderr or "")
            if re.search(r"(?<![\w-])" + re.escape(hyphen) + r"(?![\w-])", text):
                chosen = hyphen
        except (AuthToolError, CommandCancelled):
            pass
        _ESTOOL_SUBCOMMANDS[name] = chosen
        return chosen


def esptool_reset_value(value, subcommand):
    if "-" in subcommand:
        return value.replace("_", "-")
    return value.replace("-", "_")


class Manifest:
    """Validated, fingerprinted build/flasher_args.json."""

    def __init__(self, build_dir):
        self.build_dir = build_dir
        self.build_root = os.path.realpath(build_dir)
        self.path = os.path.join(build_dir, "flasher_args.json")
        self.chip = None
        self.before = None
        self.after = None
        self.stub = True
        self.write_flash_args = []
        self.entries = []
        self.fingerprint = None

    def load(self):
        if not os.path.isfile(self.path):
            raise BatchError("manifest", "flasher_args.json not found",
                             "%s (build the firmware first: idf build)" % self.path)
        try:
            with open(self.path, "rb") as handle:
                raw = handle.read()
            data = json.loads(raw.decode("utf-8"))
        except (OSError, ValueError) as exc:
            raise BatchError("manifest", "could not read flasher_args.json",
                             str(exc)) from exc

        if not isinstance(data, dict):
            raise BatchError("manifest", "flasher_args.json must be an object")
        extra = data.get("extra_esptool_args", {})
        if not isinstance(extra, dict):
            raise BatchError("manifest", "invalid extra_esptool_args",
                             "expected an object")
        unsupported = set(extra) - {"chip", "before", "after", "stub"}
        if unsupported:
            raise BatchError("manifest", "unsupported extra_esptool_args",
                             ", ".join(sorted(unsupported)))
        self.chip = extra.get("chip")
        if not isinstance(self.chip, str) or self.chip not in (
                "esp32", "esp32s2", "esp32s3", "esp32c2", "esp32c3",
                "esp32c5", "esp32c6", "esp32h2", "esp32p4"):
            raise BatchError("manifest", "invalid or unsupported chip type")
        self.before = self._validated_reset(extra.get("before", "default_reset"),
                                            "before", ("default_reset", "usb_reset",
                                                       "no_reset", "no_reset_no_sync"))
        self.after = self._validated_reset(extra.get("after", "hard_reset"),
                                           "after", ("hard_reset", "soft_reset",
                                                     "no_reset", "no_reset_stub",
                                                     "watchdog_reset"))
        self.stub = extra.get("stub", True)
        if not isinstance(self.stub, bool):
            raise BatchError("manifest", "stub must be a boolean")

        # Assign and validate the flags before inspecting them for force flags.
        self.write_flash_args = self._validated_write_flash_args(
            data.get("write_flash_args"))
        if self._uses_force_flag():
            raise BatchError("manifest", "unsupported encryption/force-flash flags",
                             " ".join(self.write_flash_args))

        self._reject_encrypted(data)
        flash_files = data.get("flash_files")
        if not isinstance(flash_files, dict) or not flash_files:
            raise BatchError("manifest", "flasher_args.json lists no flash files")
        self.entries = self._validated_entries(flash_files)
        table = data.get("partition-table")
        if not isinstance(table, dict) or not isinstance(table.get("file"), str):
            raise BatchError("manifest", "partition-table metadata is required")
        self.partition_table_offset = self._parse_offset(table.get("offset"))
        table_path = os.path.realpath(os.path.join(self.build_dir, table["file"]))
        if not any(int(offset, 16) == self.partition_table_offset and path == table_path
                   for offset, path in self.entries):
            raise BatchError("manifest", "partition-table is missing from flash_files")
        self.fingerprint = self._fingerprint(raw=raw)
        return self

    @staticmethod
    def _validated_reset(value, field, allowed):
        if not isinstance(value, str) or value.replace("-", "_") not in allowed:
            raise BatchError("manifest", "invalid %s reset setting" % field)
        return value.replace("-", "_")

    @staticmethod
    def _validated_write_flash_args(value):
        if value is None:
            return []
        if not isinstance(value, list) or not all(
                isinstance(item, str) for item in value):
            raise BatchError("manifest", "write_flash_args must be a list of strings")
        return list(value)

    @staticmethod
    def _encrypted(value):
        """Interpret every encrypted marker form safely (truthy => encrypted)."""
        if value is None or value is False:
            return False
        if isinstance(value, str):
            return value.strip().lower() not in ("", "false", "0", "no")
        if isinstance(value, (int, float)):
            return value != 0
        return bool(value)

    def _reject_encrypted(self, data):
        for name, entry in data.items():
            if isinstance(entry, dict) and self._encrypted(entry.get("encrypted")):
                raise BatchError(
                    "manifest", "encrypted image set is not supported",
                    "%s requires flash encryption provisioning" % name)

    @staticmethod
    def _parse_offset(raw_offset):
        if not isinstance(raw_offset, str) or not re.fullmatch(
                r"0[xX][0-9a-fA-F]+", raw_offset.strip()):
            raise BatchError("manifest", "invalid flash offset", repr(raw_offset))
        return int(raw_offset, 16)

    def _validated_entries(self, flash_files):
        entries = []
        seen_offsets = set()
        for raw_offset, relative in flash_files.items():
            if not isinstance(relative, str) or not relative:
                raise BatchError("manifest", "invalid flash image entry",
                                 "offset %r has no file path" % (raw_offset,))
            offset = self._parse_offset(raw_offset)
            if offset in seen_offsets:
                raise BatchError("manifest", "duplicate flash offset",
                                 raw_offset.strip())
            seen_offsets.add(offset)
            absolute = os.path.realpath(os.path.join(self.build_dir, relative))
            if os.path.commonpath([self.build_root, absolute]) != self.build_root:
                raise BatchError("manifest", "flash image escapes the build directory",
                                 "%s -> %s" % (relative, absolute))
            if not os.path.isfile(absolute):
                raise BatchError("manifest", "flash image missing", absolute)
            entries.append((offset, raw_offset.strip(), absolute))
        entries.sort(key=lambda item: item[0])
        end = 0
        for offset, _offset_text, absolute in entries:
            try:
                size = os.path.getsize(absolute)
            except OSError as exc:
                raise BatchError("manifest", "could not inspect flash image",
                                 str(exc)) from exc
            if size == 0:
                raise BatchError("manifest", "empty flash image", absolute)
            if offset < end:
                raise BatchError("manifest", "overlapping flash images", absolute)
            end = offset + size
        return [(offset_text, absolute) for _offset, offset_text, absolute in entries]

    def _uses_force_flag(self):
        forbidden = {"--force", "--ignore-flash-enc-efuse", "--encrypt",
                     "--encrypt-files"}
        return any(arg.split("=", 1)[0].replace("_", "-") in forbidden
                   for arg in self.write_flash_args)

    def _fingerprint(self, raw=None):
        try:
            if raw is None:
                with open(self.path, "rb") as handle:
                    raw = handle.read()
            digest = hashlib.sha256()
            digest.update(raw)
            digest.update(b"\x00")
            digest.update(self.chip.encode("utf-8"))
            digest.update("\x00".join(self.write_flash_args).encode("utf-8"))
            digest.update(("%s|%s|%s" % (self.before, self.after, self.stub)).encode("utf-8"))
            for offset, absolute in self.entries:
                stat = os.stat(absolute)
                digest.update(("%s|%s|%d|%d|" % (
                    offset, absolute, stat.st_size, stat.st_mtime_ns)).encode("utf-8"))
                digest.update(_file_digest(absolute).encode("utf-8"))
            return digest.hexdigest()
        except OSError as exc:
            raise BatchError("manifest", "could not fingerprint build artifacts",
                             str(exc)) from exc

    def verify_fingerprint(self):
        if self._fingerprint() != self.fingerprint:
            raise BatchError("manifest", "build artifacts changed since validation",
                             "re-run idf build and retry")

    def write_flash_command(self, port, baud):
        subcommand = esptool_subcommand("write_flash")
        command = [sys.executable, "-m", "esptool", "--chip", self.chip, "--port", port]
        if baud:
            command += ["--baud", str(baud)]
        if self.before:
            command += ["--before", esptool_reset_value(self.before, subcommand)]
        command += ["--after", esptool_reset_value("no_reset", subcommand)]
        if not self.stub:
            command += ["--no-stub"]
        command += [subcommand] + list(self.write_flash_args)
        for offset, absolute in self.entries:
            command += [offset, absolute]
        return command


def _file_digest(path):
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _file_fingerprint(path):
    stat = os.stat(path)
    return (stat.st_size, stat.st_mtime_ns, _file_digest(path))


def _fsync_dir(path):
    try:
        fd = os.open(path, os.O_RDONLY)
        try:
            os.fsync(fd)
        finally:
            os.close(fd)
    except OSError as exc:
        if exc.errno not in (errno.EINVAL, errno.ENOTSUP):
            raise


# ---------------------------------------------------------------------------
# Workbook
# ---------------------------------------------------------------------------

@dataclasses.dataclass
class CredentialRow:
    index: int
    uuid: str
    key: str
    status: object = None
    stage: object = None
    error: object = None
    port: object = None
    usb_serial: object = None
    usb_location: object = None
    device_mac: object = None
    product_key: object = None
    run_id: object = None
    updated_at: object = None
    sheet_title: object = None

    def credentials(self, product_key):
        return {"uuid": self.uuid, "auth_key": self.key, "product_key": product_key}


def _import_openpyxl():
    try:
        import openpyxl
        from openpyxl.styles import PatternFill
    except ImportError as exc:
        raise BatchError(
            "preflight", "openpyxl is not installed",
            "run idf tuya-batch-setup, or install with: %s -m pip install -r %s"
            % (sys.executable, REQUIREMENTS)) from exc
    return openpyxl, PatternFill


class Workbook:
    """Credential workbook with schema validation and atomic persistence."""

    def __init__(self, path, sheet_name=None):
        self.path = os.path.realpath(path)
        self.sheet_name = sheet_name
        self.wb = None
        self.ws = None
        self.sheet_title = None
        self.header_row = 1
        self.header = {}
        self.rows = {}
        # Allocation stays sheet-local; validation and chip bindings are global.
        self.all_rows = {}
        self._PatternFill = None
        self._mode = 0o600
        self._baseline = None

    # -- loading / validation ------------------------------------------------

    def load(self, for_write=False):
        openpyxl, pattern_fill = _import_openpyxl()
        self._PatternFill = pattern_fill
        if not os.path.isfile(self.path):
            raise BatchError("workbook", "credential workbook not found", self.path)
        self._baseline = _file_fingerprint(self.path)
        self._mode = os.stat(self.path).st_mode & 0o777
        try:
            self.wb = openpyxl.load_workbook(self.path, data_only=False)
        except Exception as exc:
            raise BatchError("workbook", "could not open workbook",
                             "%s: %s" % (self.path, exc)) from exc
        self.ws = self._select_sheet()
        self.sheet_title = self.ws.title
        self.header_row, self.header = self._find_header_row(self.ws)
        if for_write:
            self._ensure_tracking_columns()
        self.rows = {}
        self.all_rows = {}
        seen = {field: {} for field in ("uuid", "key", "device_mac")}
        for worksheet in self.wb.worksheets:
            try:
                header_row, header = self._find_header_row(worksheet)
                if header is None:
                    continue
                rows = self._parse_rows(worksheet, header_row, header)
            except BatchError as exc:
                raise BatchError(exc.stage, exc.message,
                                 "sheet %r: %s" % (worksheet.title, exc.detail or "")) from exc
            for row in rows.values():
                identity = (worksheet.title, row.index)
                for field, values in seen.items():
                    value = getattr(row, field)
                    if value is None:
                        continue
                    if value in values:
                        previous = values[value]
                        raise BatchError(
                            "workbook", "duplicate %s%s" % (
                                field, " binding" if field == "device_mac" else ""),
                            "sheet %r row %d and sheet %r row %d" % (
                                previous[0], previous[1], identity[0], identity[1]))
                    values[value] = identity
                self.all_rows[identity] = row
            if worksheet is self.ws:
                self.rows = rows
        return self

    def _find_header_row(self, worksheet):
        limit = min(worksheet.max_row or 1, 10)
        for row in range(1, limit + 1):
            mapping = {}
            duplicate = None
            for column in range(1, (worksheet.max_column or 0) + 1):
                value = worksheet.cell(row=row, column=column).value
                if not isinstance(value, str):
                    continue
                key = value.strip().lower()
                if not key:
                    continue
                if key in mapping:
                    duplicate = (key, mapping[key], column)
                else:
                    mapping[key] = column
            if not all(name in mapping for name in CREDENTIAL_HEADERS):
                continue
            if duplicate is not None:
                key, first, second = duplicate
                raise BatchError(
                    "workbook", "duplicate %s header column" % key,
                    "sheet %r row %d columns %d and %d collide after trim/case folding"
                    % (worksheet.title, row, first, second))
            return row, mapping
        return None, None

    def _select_sheet(self):
        if self.sheet_name:
            if self.sheet_name not in self.wb.sheetnames:
                raise BatchError("workbook", "worksheet not found", self.sheet_name)
            worksheet = self.wb[self.sheet_name]
            header_row, mapping = self._find_header_row(worksheet)
            if mapping is None:
                raise BatchError("workbook", "worksheet lacks uuid/key headers",
                                 self.sheet_name)
            return worksheet
        matches = []
        for name in self.wb.sheetnames:
            worksheet = self.wb[name]
            _row, mapping = self._find_header_row(worksheet)
            if mapping is not None:
                matches.append(worksheet)
        if not matches:
            raise BatchError("workbook", "no worksheet with uuid/key headers",
                             self.path)
        if len(matches) > 1:
            raise BatchError("workbook", "multiple worksheets match uuid/key headers",
                             "select one with --sheet")
        return matches[0]

    def _ensure_tracking_columns(self):
        column = self.ws.max_column or 0
        for field in TRACKING_FIELDS:
            if field not in self.header:
                column += 1
                self.ws.cell(row=self.header_row, column=column, value=field)
                self.header[field] = column

    def _credential_value(self, value, field, row):
        if value is None or value == "":
            return None
        if not isinstance(value, str) or value.startswith("="):
            raise BatchError("workbook", "credential cell is not plain text",
                             "row %d column %s (numeric/formula cells are rejected)"
                             % (row, field))
        if value.strip() != value:
            raise BatchError("workbook", "credential has surrounding whitespace",
                             "row %d column %s" % (row, field))
        if "\x00" in value:
            raise BatchError("workbook", "credential contains a NUL byte",
                             "row %d column %s" % (row, field))
        limits_field = "auth_key" if field == "key" else field
        low, high = LENGTH_LIMITS[limits_field]
        length = len(value.encode("utf-8"))
        if not low <= length <= high:
            raise BatchError("workbook", "credential UTF-8 length out of range",
                             "row %d column %s length %d not in %d..%d"
                             % (row, field, length, low, high))
        return value

    def _tracking_value(self, row, field, worksheet, header):
        column = header.get(field)
        if column is None:
            return None
        cell = worksheet.cell(row=row, column=column)
        if cell.data_type == "f":
            raise BatchError("workbook", "tracking cell is a formula",
                             "row %d column %s" % (row, field))
        return cell.value

    def _read_tracking(self, row, field, worksheet, header, allowed=None):
        value = self._tracking_value(row, field, worksheet, header)
        if value is None:
            return None
        if not isinstance(value, str):
            raise BatchError("workbook", "tracking cell is not text",
                             "row %d column %s" % (row, field))
        if value == "":
            return None
        if value != value.strip():
            raise BatchError("workbook", "tracking cell has surrounding whitespace",
                             "row %d column %s" % (row, field))
        if "\x00" in value:
            raise BatchError("workbook", "tracking cell contains a NUL byte",
                             "row %d column %s" % (row, field))
        if allowed is not None and value not in allowed:
            raise BatchError("workbook", "unrecognized %s value" % field,
                             "row %d (allowed: %s)" % (row, ", ".join(allowed)))
        return value

    @staticmethod
    def _validate_product_key(value, row):
        if value is None:
            return None
        if value.startswith("="):
            raise BatchError("workbook", "product_key is not plain text",
                             "row %d (formula cells are rejected)" % row)
        if "\x00" in value:
            raise BatchError("workbook", "product_key contains a NUL byte",
                             "row %d" % row)
        low, high = LENGTH_LIMITS["product_key"]
        length = len(value.encode("utf-8"))
        if not low <= length <= high:
            raise BatchError("workbook", "product_key UTF-8 length out of range",
                             "row %d length %d not in %d..%d"
                             % (row, length, low, high))
        return value

    def _parse_rows(self, worksheet, header_row, header):
        rows = {}
        uuid_column = header["uuid"]
        key_column = header["key"]
        seen_uuid = {}
        seen_key = {}
        seen_mac = {}
        for row in range(header_row + 1, (worksheet.max_row or 0) + 1):
            raw_uuid = worksheet.cell(row=row, column=uuid_column).value
            raw_key = worksheet.cell(row=row, column=key_column).value
            value_uuid = self._credential_value(raw_uuid, "uuid", row)
            value_key = self._credential_value(raw_key, "key", row)
            if value_uuid is None and value_key is None:
                continue
            if value_uuid is None or value_key is None:
                raise BatchError("workbook", "partial credential row",
                                 "row %d needs both uuid and key" % row)
            if value_uuid in seen_uuid:
                raise BatchError("workbook", "duplicate uuid",
                                 "rows %d and %d" % (seen_uuid[value_uuid], row))
            if value_key in seen_key:
                raise BatchError("workbook", "duplicate key",
                                 "rows %d and %d" % (seen_key[value_key], row))
            seen_uuid[value_uuid] = row
            seen_key[value_key] = row

            credential_row = CredentialRow(row, value_uuid, value_key,
                                           sheet_title=worksheet.title)
            for field in TRACKING_FIELDS:
                allowed = (VALID_STATUSES if field == "status" else
                           VALID_STAGES if field == "stage" else None)
                setattr(credential_row, field,
                        self._read_tracking(row, field, worksheet, header, allowed))
            device_mac = credential_row.device_mac
            credential_row.device_mac = normalize_mac(device_mac)
            if device_mac is not None and credential_row.device_mac is None:
                raise BatchError("workbook", "invalid device_mac",
                                 "row %d requires a 6-byte MAC address" % row)
            if credential_row.device_mac:
                if credential_row.device_mac in seen_mac:
                    raise BatchError("workbook", "duplicate device_mac binding",
                                     "rows %d and %d"
                                     % (seen_mac[credential_row.device_mac], row))
                seen_mac[credential_row.device_mac] = row
            credential_row.product_key = self._validate_product_key(
                credential_row.product_key, row)
            self._validate_lifecycle(credential_row)
            rows[row] = credential_row
        return rows

    @staticmethod
    def _validate_lifecycle(row):
        tracked = ("stage", "error", "port", "usb_serial", "usb_location",
                   "device_mac", "product_key", "run_id", "updated_at")
        if row.status in (None, STATUS_UNUSED):
            present = [field for field in tracked if getattr(row, field) is not None]
            if present:
                raise BatchError(
                    "workbook", "unused credential row has tracking data",
                    "row %d fields: %s" % (row.index, ", ".join(present)))
            return

        required = ("stage", "device_mac", "product_key", "port", "run_id",
                    "updated_at")
        missing = [field for field in required if getattr(row, field) is None]
        if missing:
            raise BatchError(
                "workbook", "%s credential row is incomplete" % row.status,
                "row %d missing: %s" % (row.index, ", ".join(missing)))
        if row.status == STATUS_USED:
            if row.stage != STAGE_COMPLETE or row.error is not None:
                raise BatchError(
                    "workbook", "used credential row has invalid lifecycle",
                    "row %d requires stage complete and no error" % row.index)
        elif row.status == STATUS_FAIL:
            if row.stage == STAGE_COMPLETE or row.error is None:
                raise BatchError(
                    "workbook", "failed credential row has invalid lifecycle",
                    "row %d requires a non-complete stage and error" % row.index)
        elif row.stage == STAGE_COMPLETE or row.error is not None:
            raise BatchError(
                "workbook", "in_progress credential row has invalid lifecycle",
                "row %d requires a non-complete stage and no error" % row.index)

    def find_row_by_uuid(self, value):
        for row in self.rows.values():
            if row.uuid == value:
                return row
        return None

    # -- mutation / persistence ---------------------------------------------

    def fill_for(self, status):
        if status == STATUS_USED:
            color = _FILL_GREEN
        elif status == STATUS_FAIL:
            color = _FILL_RED
        else:
            color = _FILL_NEUTRAL
        if color is None:
            return self._PatternFill(fill_type=None)
        return self._PatternFill(start_color=color, end_color=color, fill_type="solid")

    def set_tracking(self, row, field, value, fill=None):
        column = self.header.get(field)
        if column is None:
            raise PersistenceError("tracking column missing", field)
        cell = self.ws.cell(row=row, column=column)
        cell.value = value
        if isinstance(value, str):
            cell.data_type = "s"
        if fill is not None:
            cell.fill = fill

    def save(self):
        self._check_external_edit()
        directory = os.path.dirname(self.path) or "."
        fd, temp_path = tempfile.mkstemp(prefix=".batch_", suffix=".xlsx",
                                         dir=directory)
        os.close(fd)
        try:
            self.wb.save(temp_path)
            with open(temp_path, "rb+") as handle:
                os.fsync(handle.fileno())
            os.chmod(temp_path, self._mode)
            self._check_external_edit()
            os.replace(temp_path, self.path)
            _fsync_dir(directory)
        except Exception:
            try:
                os.unlink(temp_path)
            except OSError:
                pass
            raise
        self._baseline = _file_fingerprint(self.path)

    def _check_external_edit(self):
        try:
            current = _file_fingerprint(self.path)
        except OSError as exc:
            raise PersistenceError("workbook disappeared during the run",
                                   str(exc)) from exc
        if current != self._baseline:
            raise PersistenceError(
                "workbook changed on disk since it was opened; refusing to overwrite",
                "close other editors and re-run")


def allocate_rows(rows, devices):
    """Deterministically pair sheet-order unused rows with devices."""
    eligible = [row for row in rows.values()
                if row.status in (None, STATUS_UNUSED) and not row.device_mac]
    eligible.sort(key=lambda row: row.index)
    if len(eligible) < len(devices):
        raise BatchError("workbook", "not enough unused credential rows",
                         "need %d, have %d" % (len(devices), len(eligible)))
    return list(zip(devices, eligible[:len(devices)]))


# ---------------------------------------------------------------------------
# Lock, backup and journal
# ---------------------------------------------------------------------------

class WorkbookLock:
    """Exclusive advisory lock on a stable companion file."""

    def __init__(self, xlsx_path):
        self.path = os.path.realpath(xlsx_path) + ".lock"
        self._fd = None

    def acquire(self):
        if fcntl is None:
            raise BatchError("preflight", "advisory file locking is unavailable",
                             "batch flashing requires macOS or Linux")
        try:
            self._fd = os.open(self.path, os.O_CREAT | os.O_RDWR, 0o600)
        except OSError as exc:
            raise BatchError("preflight", "could not create the workbook lock",
                             str(exc)) from exc
        try:
            fcntl.flock(self._fd, fcntl.LOCK_EX | fcntl.LOCK_NB)
        except OSError as exc:
            os.close(self._fd)
            self._fd = None
            raise BatchError("preflight", "another batch run is using this workbook",
                             self.path) from exc

    def release(self):
        if self._fd is not None:
            try:
                fcntl.flock(self._fd, fcntl.LOCK_UN)
            except OSError:
                pass
            os.close(self._fd)
            self._fd = None


def write_backup(xlsx_path):
    """Copy the workbook to a protected pre-run backup beside it."""
    xlsx_path = os.path.realpath(xlsx_path)
    directory = os.path.dirname(xlsx_path) or "."
    name = os.path.basename(xlsx_path)
    target = os.path.join(directory, name + ".batch-backup.xlsx")
    fd, temp_path = tempfile.mkstemp(prefix=name + ".", suffix=".bak", dir=directory)
    os.close(fd)
    try:
        shutil.copyfile(xlsx_path, temp_path)
        os.chmod(temp_path, 0o600)
        with open(temp_path, "rb+") as handle:
            os.fsync(handle.fileno())
        os.replace(temp_path, target)
        _fsync_dir(directory)
    except OSError as exc:
        try:
            os.unlink(temp_path)
        except OSError:
            pass
        raise PersistenceError("could not write the pre-run backup", str(exc)) from exc
    return target


def verify_workbook_directory_writable(xlsx_path):
    """Verify atomic workbook persistence without changing the workbook."""
    xlsx_path = os.path.realpath(xlsx_path)
    directory = os.path.dirname(xlsx_path) or "."
    name = os.path.basename(xlsx_path)
    temp_path = None
    try:
        fd, temp_path = tempfile.mkstemp(prefix=name + ".", suffix=".write-test",
                                         dir=directory)
        with os.fdopen(fd, "wb") as handle:
            handle.write(b"batch-flash-write-test\n")
            handle.flush()
            os.fsync(handle.fileno())
        os.unlink(temp_path)
        temp_path = None
        _fsync_dir(directory)
    except OSError as exc:
        if temp_path is not None:
            try:
                os.unlink(temp_path)
            except OSError:
                pass
        raise PersistenceError("workbook directory is not writable", str(exc)) from exc


class Journal:
    """Secret-free append-only JSONL journal, fsynced after each record."""

    def __init__(self, xlsx_path):
        self.path = os.path.realpath(xlsx_path) + ".batch-journal.jsonl"
        try:
            fd = os.open(self.path, os.O_WRONLY | os.O_CREAT | os.O_APPEND, 0o600)
        except OSError as exc:
            raise PersistenceError("could not open the run journal", str(exc)) from exc
        self._file = os.fdopen(fd, "a", encoding="utf-8")

    def append(self, event, **fields):
        record = {"ts": _utc_now(), "event": event}
        record.update(fields)
        self._file.write(json.dumps(record, sort_keys=True, default=str) + "\n")
        self._file.flush()
        os.fsync(self._file.fileno())

    def close(self):
        self._file.close()


# ---------------------------------------------------------------------------
# Planning and probing
# ---------------------------------------------------------------------------

@dataclasses.dataclass
class Assignment:
    candidate: PortCandidate
    row: CredentialRow
    pid: str
    retry_mac: object = None
    mac: object = None


@dataclasses.dataclass
class ProbeResult:
    chip: object = None
    mac: object = None
    error: object = None


@dataclasses.dataclass
class WorkResult:
    path: str
    row: int
    result: str
    stage: str
    detail: str = ""


def normalize_mac(value):
    if not value or not re.fullmatch(
            r"(?:[0-9a-fA-F]{12}|[0-9a-fA-F]{2}([:-])"
            r"[0-9a-fA-F]{2}(?:\1[0-9a-fA-F]{2}){4}|"
            r"[0-9a-fA-F]{4}(?:\.[0-9a-fA-F]{4}){2})", value):
        return None
    hex_only = re.sub(r"[:.-]", "", value).lower()
    return ":".join(hex_only[i:i + 2] for i in range(0, 12, 2))


# Exact package aliases only: a shared prefix does not establish chip family.
_CHIP_FAMILY_ALIASES = {
    "esp32": (
        "esp32d0wd", "esp32d0wdq6", "esp32d0wdv3", "esp32d0wdq6v3",
        "esp32d2wd", "esp32s0wd", "esp32u4wdh", "esp32picod4",
        "esp32picov302", "esp32picov3"),
    "esp32c3": ("esp8685",),
    "esp32c6": ("esp32c6fh4", "esp32c6fh8"),
}


def normalize_chip(value):
    normalized = re.sub(r"[^a-z0-9]", "", (value or "").lower())
    for family, aliases in _CHIP_FAMILY_ALIASES.items():
        if normalized in aliases:
            return family
    return normalized


def chips_compatible(expected, reported):
    expected = normalize_chip(expected)
    return bool(expected) and expected == normalize_chip(reported)


def _require_pid(args):
    if not args.pid:
        raise BatchError("preflight", "--pid is required", "supply the Tuya product key")
    if args.pid.startswith("="):
        raise BatchError("preflight", "--pid must not start with '='")
    if args.pid != args.pid.strip():
        raise BatchError("preflight", "--pid has surrounding whitespace")
    if "\x00" in args.pid:
        raise BatchError("preflight", "--pid contains a NUL byte")
    low, high = LENGTH_LIMITS["product_key"]
    length = len(args.pid.encode("utf-8"))
    if not low <= length <= high:
        raise BatchError("preflight", "--pid UTF-8 length out of range",
                         "%d not in %d..%d" % (length, low, high))


def _validate_idf_tools():
    try:
        find_idf_tools_structured()
    except AuthToolError as exc:
        raise BatchError(exc.stage, exc.message, exc.detail) from exc
    for module in ("esptool", "esp_idf_nvs_partition_gen"):
        if not _module_available(module):
            raise BatchError("preflight", "%s is not importable" % module,
                             "run idf tuya-batch-setup and use the ESP-IDF Python environment")


def _module_available(name):
    import importlib.util
    return importlib.util.find_spec(name) is not None


def _validate_retry_structure(args):
    if len(args.device) != 1:
        raise BatchError("preflight",
                         "--retry-uuid requires exactly one --device PORT",
                         "supplied %d device path(s)" % len(args.device))


def plan_assignments(args, workbook, candidates):
    interrupted = [str(row.index) for row in workbook.rows.values()
                   if row.status == STATUS_IN_PROGRESS]
    if interrupted:
        print("WARNING: interrupted in_progress rows %s remain reserved; "
              "use explicit --retry-uuid with the original device, not a new row."
              % ", ".join(interrupted), file=sys.stderr)
    if args.retry_uuid:
        row = workbook.find_row_by_uuid(args.retry_uuid)
        if row is None:
            raise BatchError("workbook", "no credential row with the retry uuid")
        if row.status == STATUS_USED:
            raise BatchError("workbook", "used rows cannot be retried",
                             "row %d is terminal" % row.index)
        if row.status not in (STATUS_FAIL, STATUS_IN_PROGRESS):
            raise BatchError("workbook", "row is not retryable",
                             "row %d status %r (only fail/in_progress)"
                             % (row.index, row.status))
        if not row.product_key:
            raise BatchError("workbook", "row has no recorded product_key",
                             "manual review required for row %d" % row.index)
        if row.product_key != args.pid:
            raise BatchError("preflight", "retry requires the original --pid",
                             "row %d recorded a different product key" % row.index)
        if not row.device_mac:
            raise BatchError("workbook", "row has no recorded chip identity",
                             "manual review required for row %d" % row.index)
        return [Assignment(candidates[0], row, args.pid, retry_mac=row.device_mac)]
    pairs = allocate_rows(workbook.rows, candidates)
    return [Assignment(candidate, row, args.pid) for candidate, row in pairs]


def probe_device(candidate, manifest, baud, timeout, cancel_event):
    # Conservative preflight: pin the manifest chip and never reset into the
    # application before credentials are written.
    subcommand = esptool_subcommand("read_mac")
    command = [sys.executable, "-m", "esptool", "--chip", manifest.chip,
               "--port", candidate.path, "--after",
               esptool_reset_value("no_reset", subcommand)]
    if baud:
        command += ["--baud", str(baud)]
    if manifest.before:
        command += ["--before", esptool_reset_value(manifest.before, subcommand)]
    if not manifest.stub:
        command += ["--no-stub"]
    command += [subcommand]
    try:
        result = run_command(command, "probe", timeout=timeout,
                             cancel_event=cancel_event)
    except CommandCancelled:
        raise
    except AuthToolError as exc:
        return ProbeResult(error=exc.detail or exc.message)
    output = (result.stdout or "") + "\n" + (result.stderr or "")
    mac_match = _BASE_MAC_RE.search(output) or _MAC_RE.search(output)
    chip_match = _CHIP_RE.search(output)
    if not mac_match:
        return ProbeResult(error="could not read chip MAC (unexpected esptool output)")
    return ProbeResult(chip=chip_match.group(1) if chip_match else None,
                       mac=normalize_mac(mac_match.group(1)))


def finalize_device(assignment, manifest, baud, timeout, cancel_event):
    """Apply the manifest reset only after exact credential read-back verification."""
    check_cancelled(cancel_event, STAGE_AUTH_VERIFY)
    subcommand = esptool_subcommand("read_mac")
    command = [sys.executable, "-m", "esptool", "--chip", manifest.chip,
               "--port", assignment.candidate.path]
    if baud:
        command += ["--baud", str(baud)]
    if manifest.before:
        command += ["--before", esptool_reset_value(manifest.before, subcommand)]
    command += ["--after", esptool_reset_value(manifest.after or "hard_reset", subcommand)]
    if not manifest.stub:
        command += ["--no-stub"]
    command += [subcommand]
    try:
        verify_candidates_frozen([assignment.candidate])
    except BatchError as exc:
        raise BatchError(STAGE_AUTH_VERIFY, exc.message, exc.detail, exit_code=1) from exc
    return run_command(command, STAGE_AUTH_VERIFY, timeout=timeout,
                       cancel_event=cancel_event)


def verify_assignment_probe(assignment, manifest, baud, timeout, cancel_event,
                            stage):
    result = probe_device(assignment.candidate, manifest, baud,
                          min(timeout, PROBE_TIMEOUT_CAP), cancel_event)
    if result.error:
        raise BatchError(stage, "device identity re-probe failed", result.error,
                         exit_code=1)
    if result.mac != assignment.mac:
        raise BatchError(
            stage, "device identity changed after reservation",
            "%s reports %s, expected %s"
            % (assignment.candidate.path, result.mac or "unknown", assignment.mac),
            exit_code=1)
    if not chips_compatible(manifest.chip, result.chip):
        raise BatchError(
            stage, "device chip changed after reservation",
            "%s reports %s, build targets %s"
            % (assignment.candidate.path, result.chip or "unknown", manifest.chip),
            exit_code=1)


def probe_barrier(candidates, args, manifest, cancel_event, progress=None):
    jobs = args.jobs or len(candidates)
    jobs = max(1, min(jobs, len(candidates)))
    timeout = min(args.timeout, PROBE_TIMEOUT_CAP)
    results = {}
    with concurrent.futures.ThreadPoolExecutor(max_workers=jobs) as pool:
        futures = {}
        for candidate in candidates:
            def task(candidate=candidate):
                if progress is not None:
                    progress.probe_start(candidate.path)
                return probe_device(candidate, manifest, args.baud, timeout,
                                    cancel_event)
            futures[pool.submit(task)] = candidate
        pending = set(futures)
        while pending:
            done, pending = concurrent.futures.wait(
                pending, timeout=_PROGRESS_TICK,
                return_when=concurrent.futures.FIRST_COMPLETED)
            for future in done:
                candidate = futures[future]
                try:
                    result = future.result()
                except CommandCancelled:
                    result = ProbeResult(error="cancelled")
                results[candidate.path] = result
                if progress is not None:
                    progress.probe(candidate.path, result.error is None)
            if progress is not None:
                if cancel_event.is_set():
                    progress.set_stopping()
                progress.tick()
    return results


def validate_probes(probe_results, manifest, workbook, args):
    failures = ["%s - %s" % (path, result.error)
                for path, result in probe_results.items() if result.error]
    if failures:
        raise BatchError("probe", "device preflight failed; no credentials consumed",
                         "; ".join(sorted(failures)))

    chips = {}
    for path, result in probe_results.items():
        chips.setdefault(result.mac, []).append(path)
    duplicates = {mac: paths for mac, paths in chips.items() if len(paths) > 1}
    if duplicates:
        detail = "; ".join("%s on %s" % (mac, ", ".join(sorted(paths)))
                           for mac, paths in sorted(duplicates.items()))
        raise BatchError("probe", "duplicate chip detected across ports", detail)

    for path, result in probe_results.items():
        if not chips_compatible(manifest.chip, result.chip):
            raise BatchError("probe", "device chip does not match the build",
                             "%s reports %s, build targets %s"
                             % (path, result.chip or "unknown", manifest.chip))

    retry_row = None
    if args.retry_uuid:
        retry_row = workbook.find_row_by_uuid(args.retry_uuid)
    bound = {}
    for row in workbook.all_rows.values():
        if row.device_mac:
            bound[row.device_mac] = row
    for path, result in probe_results.items():
        row = bound.get(result.mac)
        if row is not None and (retry_row is None or
                (row.sheet_title, row.index) != (retry_row.sheet_title, retry_row.index)):
            raise BatchError("probe", "device is already bound to a workbook row",
                             "%s (chip %s) is sheet %r row %d; disconnect it or use explicit retry"
                             % (path, result.mac, row.sheet_title, row.index))
    if retry_row is not None:
        result = probe_results[args.device[0]]
        if result.mac != retry_row.device_mac:
            raise BatchError("probe", "retry device chip does not match the record",
                             "%s reports %s, row %d recorded %s"
                             % (args.device[0], result.mac, retry_row.index,
                                retry_row.device_mac))


# ---------------------------------------------------------------------------
# Coordinator and workers
# ---------------------------------------------------------------------------

class Coordinator:
    """Serializes journal and workbook writes; signals persistence failures."""

    def __init__(self, workbook, journal, run_id, cancel_event, progress=None):
        self.workbook = workbook
        self.journal = journal
        self.run_id = run_id
        self.cancel_event = cancel_event
        self.progress = progress
        self.failure = None
        self._lock = threading.RLock()
        self._secrets = set()

    def register_secrets(self, assignments):
        for assignment in assignments:
            self._secrets.add(assignment.row.key)

    def sanitize(self, text):
        if not text:
            return text
        for secret in self._secrets:
            if secret:
                text = text.replace(secret, "[redacted]")
        return text

    def record_persistence_failure(self, error):
        with self._lock:
            if self.failure is None:
                if isinstance(error, PersistenceError):
                    self.failure = PersistenceError(
                        self.sanitize(error.message), self.sanitize(error.detail))
                else:
                    self.failure = PersistenceError(self.sanitize(str(error)))
            self.cancel_event.set()
            return self.failure

    def reserve(self, assignments, pid):
        with self._lock:
            now = _utc_now()
            try:
                for assignment in assignments:
                    self.journal.append(
                        "reserve", run_id=self.run_id,
                        sheet=self.workbook.sheet_title, row=assignment.row.index,
                        uuid=assignment.row.uuid, stage=STAGE_RESERVED,
                        port=assignment.candidate.path,
                        usb_serial=assignment.candidate.serial_number,
                        usb_location=assignment.candidate.location,
                        device_mac=assignment.mac, product_key=pid)
                for assignment in assignments:
                    row = assignment.row
                    self.workbook.set_tracking(row.index, "status", STATUS_IN_PROGRESS,
                                               fill=self.workbook.fill_for(STATUS_IN_PROGRESS))
                    self.workbook.set_tracking(row.index, "stage", STAGE_RESERVED)
                    self.workbook.set_tracking(row.index, "error", None)
                    self.workbook.set_tracking(row.index, "port", assignment.candidate.path)
                    self.workbook.set_tracking(row.index, "usb_serial",
                                               assignment.candidate.serial_number)
                    self.workbook.set_tracking(row.index, "usb_location",
                                               assignment.candidate.location)
                    self.workbook.set_tracking(row.index, "device_mac", assignment.mac)
                    self.workbook.set_tracking(row.index, "product_key", pid)
                    self.workbook.set_tracking(row.index, "run_id", self.run_id)
                    self.workbook.set_tracking(row.index, "updated_at", now)
                    row.status = STATUS_IN_PROGRESS
                    row.stage = STAGE_RESERVED
                    row.error = None
                    row.port = assignment.candidate.path
                    row.usb_serial = assignment.candidate.serial_number
                    row.usb_location = assignment.candidate.location
                    row.device_mac = assignment.mac
                    row.product_key = pid
                    row.run_id = self.run_id
                    row.updated_at = now
                self.workbook.save()
            except Exception as exc:
                raise self.record_persistence_failure(exc) from exc

    def set_stage(self, assignment, stage):
        with self._lock:
            if self.failure is not None:
                raise self.failure
            try:
                self.journal.append(
                    "stage", run_id=self.run_id, sheet=self.workbook.sheet_title,
                    row=assignment.row.index, uuid=assignment.row.uuid, stage=stage)
            except Exception as exc:
                raise self.record_persistence_failure(exc) from exc
            assignment.row.stage = stage
            if self.progress is not None:
                self.progress.stage(assignment.candidate.path, stage)
            else:
                print_progress(assignment.candidate.path, stage, assignment.row.index)

    def finish(self, assignment, status, stage, error):
        with self._lock:
            if self.failure is not None:
                raise self.failure
            now = _utc_now()
            detail = self.sanitize(error)
            if detail:
                detail = detail[:ERROR_MAX_LEN]
            try:
                self.journal.append(
                    "outcome", run_id=self.run_id, sheet=self.workbook.sheet_title,
                    row=assignment.row.index, uuid=assignment.row.uuid,
                    status=status, stage=stage, error=detail,
                    port=assignment.candidate.path, device_mac=assignment.mac)
                row = assignment.row
                self.workbook.set_tracking(row.index, "status", status,
                                           fill=self.workbook.fill_for(status))
                self.workbook.set_tracking(row.index, "stage", stage)
                self.workbook.set_tracking(row.index, "error", detail)
                self.workbook.set_tracking(row.index, "port", assignment.candidate.path)
                self.workbook.set_tracking(row.index, "device_mac", assignment.mac)
                self.workbook.set_tracking(row.index, "updated_at", now)
                row.status = status
                row.stage = stage
                row.error = detail
                row.port = assignment.candidate.path
                row.device_mac = assignment.mac
                row.updated_at = now
                self.workbook.save()
            except Exception as exc:
                raise self.record_persistence_failure(exc) from exc


def flash_device(assignment, manifest, coordinator, baud, timeout):
    try:
        return _flash_device(assignment, manifest, coordinator, baud, timeout)
    except PersistenceError as exc:
        stage = _current_failure_stage(assignment)
        return WorkResult(assignment.candidate.path, assignment.row.index,
                          RESULT_NOT_SAVED, stage,
                          str(coordinator.record_persistence_failure(exc)))
    except Exception as exc:  # pragma: no cover - defensive isolation
        return _record_unexpected_failure(coordinator, assignment, exc)


def _flash_device(assignment, manifest, coordinator, baud, timeout):
    path = assignment.candidate.path
    row = assignment.row.index
    # Queued workers must do nothing once the batch is aborting.
    if coordinator.cancel_event.is_set() or coordinator.failure is not None:
        return WorkResult(path, row, RESULT_CANCELLED, STAGE_FIRMWARE, "cancelled")
    try:
        manifest.verify_fingerprint()
        coordinator.set_stage(assignment, STAGE_FIRMWARE)
        check_cancelled(coordinator.cancel_event, STAGE_FIRMWARE)
        flash_command = manifest.write_flash_command(path, baud)
        verify_candidates_frozen([assignment.candidate])
        verify_assignment_probe(assignment, manifest, baud, timeout,
                                coordinator.cancel_event, STAGE_FIRMWARE)
        run_command(flash_command, STAGE_FIRMWARE, timeout=timeout,
                    cancel_event=coordinator.cancel_event)
    except CommandCancelled:
        return WorkResult(path, row, RESULT_CANCELLED, STAGE_FIRMWARE, "cancelled")
    except (AuthToolError, BatchError) as exc:
        return _record_failure(coordinator, assignment, STAGE_FIRMWARE, exc)

    try:
        coordinator.set_stage(assignment, STAGE_AUTH_WRITE)
        check_cancelled(coordinator.cancel_event, STAGE_AUTH_WRITE)
        verify_candidates_frozen([assignment.candidate])
        verify_assignment_probe(assignment, manifest, baud, timeout,
                                coordinator.cancel_event, STAGE_AUTH_WRITE)
        write_verify_identity(
            path, baud, assignment.row.credentials(assignment.pid),
            timeout=timeout, cancel_event=coordinator.cancel_event,
            stage_callback=lambda stage: coordinator.set_stage(assignment, stage),
            before=manifest.before, after="no_reset",
            partition_table_offset=manifest.partition_table_offset)
    except CommandCancelled:
        return WorkResult(path, row, RESULT_CANCELLED,
                          _current_failure_stage(assignment), "cancelled")
    except BatchError as exc:
        return _record_failure(coordinator, assignment, STAGE_AUTH_WRITE, exc)
    except AuthToolError as exc:
        return _record_failure(coordinator, assignment, exc.stage, exc)

    try:
        coordinator.set_stage(assignment, STAGE_AUTH_VERIFY)
        finalize_device(assignment, manifest, baud, timeout, coordinator.cancel_event)
    except CommandCancelled:
        return WorkResult(path, row, RESULT_CANCELLED, STAGE_AUTH_VERIFY, "cancelled")
    except (AuthToolError, BatchError) as exc:
        return _record_failure(coordinator, assignment, STAGE_AUTH_VERIFY, exc)

    coordinator.finish(assignment, STATUS_USED, STAGE_COMPLETE, None)
    return WorkResult(path, row, RESULT_PASS, STAGE_COMPLETE, "verified; saved as used")


def _current_failure_stage(assignment):
    stage = assignment.row.stage
    if stage in VALID_STAGES and stage != STAGE_COMPLETE:
        return stage
    return STAGE_FIRMWARE


def _record_failure(coordinator, assignment, stage, exc):
    detail = "%s: %s" % (exc.message, exc.detail) if exc.detail else exc.message
    detail = coordinator.sanitize(detail)
    try:
        coordinator.finish(assignment, STATUS_FAIL, stage, detail)
    except PersistenceError as persistence_error:
        return WorkResult(assignment.candidate.path, assignment.row.index,
                          RESULT_NOT_SAVED, stage,
                          str(coordinator.record_persistence_failure(persistence_error)))
    return WorkResult(assignment.candidate.path, assignment.row.index,
                      RESULT_FAIL, stage, detail)


def _record_unexpected_failure(coordinator, assignment, exc):
    stage = _current_failure_stage(assignment)
    detail = coordinator.sanitize("%s: %s" % (type(exc).__name__, exc))
    try:
        coordinator.finish(assignment, STATUS_FAIL, stage, detail)
    except Exception as persistence_error:
        return WorkResult(
            assignment.candidate.path, assignment.row.index, RESULT_NOT_SAVED, stage,
            str(coordinator.record_persistence_failure(persistence_error)))
    return WorkResult(assignment.candidate.path, assignment.row.index,
                      RESULT_FAIL, stage, detail)


def run_workers(assignments, manifest, coordinator, jobs, baud, timeout):
    results = {}
    progress = coordinator.progress
    workers = max(1, min(jobs, len(assignments)))
    with concurrent.futures.ThreadPoolExecutor(max_workers=workers) as pool:
        futures = {pool.submit(flash_device, assignment, manifest, coordinator,
                               baud, timeout): assignment for assignment in assignments}
        pending = set(futures)
        while pending:
            done, pending = concurrent.futures.wait(
                pending, timeout=_PROGRESS_TICK,
                return_when=concurrent.futures.FIRST_COMPLETED)
            for future in done:
                assignment = futures[future]
                try:
                    result = future.result()
                except concurrent.futures.CancelledError:
                    result = WorkResult(assignment.candidate.path, assignment.row.index,
                                        RESULT_CANCELLED,
                                        assignment.row.stage or STAGE_FIRMWARE, "cancelled")
                except PersistenceError as exc:
                    # A persistence failure is batch-wide: cancel peers, do not
                    # report it as an ordinary device result.
                    result = WorkResult(assignment.candidate.path, assignment.row.index,
                                        RESULT_NOT_SAVED,
                                        assignment.row.stage or STAGE_FIRMWARE,
                                        str(coordinator.record_persistence_failure(exc)))
                except Exception as exc:  # pragma: no cover - defensive isolation
                    result = _record_unexpected_failure(coordinator, assignment, exc)
                if result.result == RESULT_NOT_SAVED:
                    result.detail = str(coordinator.record_persistence_failure(
                        PersistenceError(result.detail or RESULT_NOT_SAVED)))
                results[assignment.candidate.path] = result
                if progress is not None:
                    progress.finish(result.path, result.row, result.result)
                else:
                    print_progress(result.path, result.result, result.row)
            # Cancellation is batch-wide: handle it on every poll, including
            # timeouts with no completed future, so queued peers are dropped
            # and the operator sees the stopping notice while slow commands
            # are still being terminated.
            if coordinator.failure is not None and not coordinator.cancel_event.is_set():
                coordinator.cancel_event.set()
            if coordinator.cancel_event.is_set():
                if progress is not None:
                    progress.set_stopping()
                for pending_future in pending:
                    pending_future.cancel()
            if progress is not None:
                progress.tick()
    return results


# ---------------------------------------------------------------------------
# Reporting
# ---------------------------------------------------------------------------

_PRINT_LOCK = threading.Lock()

_PROGRESS_TICK = 0.25
_MAX_LIVE_ROWS = 10
_MIN_LIVE_WIDTH = 60
_MIN_LIVE_HEIGHT = 6

STAGE_LABELS = {
    STAGE_RESERVED: "reserved",
    STAGE_FIRMWARE: "firmware / device checks",
    STAGE_AUTH_WRITE: "writing credentials",
    STAGE_AUTH_READ: "reading back credentials",
    STAGE_AUTH_VERIFY: "verifying / final reset",
    STAGE_COMPLETE: "complete",
    RESULT_PASS: "PASS (verified)",
    RESULT_FAIL: "FAIL",
    RESULT_CANCELLED: "CANCELLED",
    RESULT_NOT_SAVED: "RESULT NOT SAVED",
}

_PROBE_STATE_LABELS = {
    "waiting": "waiting",
    "probing": "probing",
    "response": "response received",
    "error": "probe error",
}


def print_progress(path, stage, row):
    with _PRINT_LOCK:
        print("[%s] %s (row %d)" % (stage, path, row), flush=True)


def _format_duration(seconds):
    if seconds is None or seconds < 0:
        return "--:--"
    total = int(seconds)
    hours, remainder = divmod(total, 3600)
    minutes, secs = divmod(remainder, 60)
    if hours:
        return "%d:%02d:%02d" % (hours, minutes, secs)
    return "%02d:%02d" % (minutes, secs)


def _sanitize_text(text):
    return "".join(char if 0x20 <= ord(char) < 0x7f else "?" for char in str(text))


def _truncate(text, width):
    if width <= 0:
        return ""
    if len(text) <= width:
        return text
    if width <= 3:
        return text[:width]
    return text[:width - 3] + "..."


def _shorten_path(path, width):
    if width <= 0:
        return ""
    if len(path) <= width:
        return path
    if width <= 3:
        return path[-width:]
    return "..." + path[-(width - 3):]


def _is_tty(stream):
    try:
        return bool(stream.isatty())
    except (AttributeError, ValueError):
        return False


def _color_enabled(stream):
    if os.environ.get("NO_COLOR"):
        return False
    if os.environ.get("TERM") == "dumb":
        return False
    return _is_tty(stream)


def _terminal_size(stream):
    try:
        size = os.get_terminal_size(stream.fileno())
        return size.columns, size.lines
    except (AttributeError, ValueError, OSError):
        pass
    try:
        return int(os.environ.get("COLUMNS", "80")), int(os.environ.get("LINES", "24"))
    except (TypeError, ValueError):
        return 80, 24


def _live_output_supported(stream, size):
    if not _is_tty(stream) or os.environ.get("TERM") == "dumb":
        return False
    columns, lines = size if size is not None else _terminal_size(stream)
    return columns >= _MIN_LIVE_WIDTH and lines >= _MIN_LIVE_HEIGHT


@dataclasses.dataclass
class _ProgressEntry:
    path: str
    row: int
    probe_state: str = "waiting"
    probe_started: object = None
    probe_done: object = None
    stage: object = None
    stage_started: object = None
    result: object = None
    finished_at: object = None
    device_started: object = None


class ProgressDisplay:
    """Read-only batch progress: a live TTY table or plain event lines.

    Workers only record state; the collector threads call ``tick`` to refresh
    the live frame. The display never reserves rows, classifies probes or
    declares success: terminal outcomes come solely from collected work
    results.
    """

    def __init__(self, assignments, stream=None, clock=time.monotonic,
                 size=None, live=None, interval=_PROGRESS_TICK):
        self._stream = stream if stream is not None else sys.stdout
        self._clock = clock
        self._interval = interval
        self._size = size
        self._entries = {}
        self._order = []
        for assignment in assignments:
            entry = _ProgressEntry(assignment.candidate.path, assignment.row.index)
            self._entries[entry.path] = entry
            self._order.append(entry.path)
        self._total = len(self._order)
        self._lock = threading.Lock()
        # Plain lines recorded by workers, flushed by the collector thread so
        # no worker ever performs console I/O (it must not stall journaling).
        self._pending = []
        self._phase = None
        self._started = False
        self._stopping = False
        self._run_started = None
        self._program_started = None
        self._last_render = float("-inf")
        self._rendered = 0
        self._width = 80
        self._live = self._resolve_live(size, live)
        self._use_color = _color_enabled(self._stream)

    def _resolve_live(self, size, live):
        supported = (_live_output_supported(self._stream, size) if live is None
                     else bool(live))
        if not supported:
            return False
        columns, lines = size if size is not None else _terminal_size(self._stream)
        self._width = columns
        # Never hide devices: fall back to plain lines when the batch does not fit.
        return self._geometry_fits(columns, lines)

    # -- lifecycle -----------------------------------------------------------

    def start(self, phase):
        now = self._clock()
        with self._lock:
            self._phase = phase
            self._started = True
            self._last_render = float("-inf")
            if self._run_started is None:
                self._run_started = now
            if phase == "flashing" and self._program_started is None:
                self._program_started = now
                for path in self._order:
                    if self._entries[path].result is None:
                        self._entries[path].device_started = now
            live = self._live
        if not live and phase == "probing":
            self._print_plain("Probing %d device(s)..." % self._total)

    def stop(self, summarize=True):
        with self._lock:
            if not self._started:
                return
            self._started = False
            phase = self._phase
            live = self._live
            if live and self._size is None:
                columns, lines_count = _terminal_size(self._stream)
                if self._geometry_fits(columns, lines_count):
                    self._width = columns
                else:
                    # The terminal shrank below a usable size; leave the last
                    # frame in place instead of drawing over unknown geometry.
                    self._live = False
                    live = False
            if live:
                now = self._clock()
                self._last_render = now
                lines = self._build_lines(now)
            pending = self._pending
            self._pending = []
        if live:
            with _PRINT_LOCK:
                if self._write_frame(lines):
                    self._write("\n")
                self._rendered = 0
                self._write_pending_locked(pending)
        else:
            with _PRINT_LOCK:
                if self._rendered:
                    self._write("\n")
                    self._rendered = 0
                self._write_pending_locked(pending)
        if summarize:
            self._print_plain_summary(phase)

    # -- event producers (worker threads) ------------------------------------

    def probe_start(self, path):
        now = self._clock()
        with self._lock:
            entry = self._entries.get(path)
            if entry is None or entry.probe_state != "waiting":
                return
            entry.probe_state = "probing"
            entry.probe_started = now
            row = entry.row
            if not self._live:
                self._pending.append(
                    "[probe] %s (row %d) | probing started" % (path, row))

    def probe(self, path, ok):
        now = self._clock()
        with self._lock:
            entry = self._entries.get(path)
            if entry is None or entry.probe_state in ("response", "error"):
                return
            entry.probe_state = "response" if ok else "error"
            entry.probe_done = now
            row = entry.row
            if not self._live:
                label = "response received" if ok else "probe error"
                self._pending.append("[probe] %s (row %d) | %s" % (path, row, label))

    def stage(self, path, stage):
        now = self._clock()
        with self._lock:
            entry = self._entries.get(path)
            if entry is None or entry.result is not None or entry.stage == stage:
                return
            entry.stage = stage
            entry.stage_started = now
            if not self._live:
                label = STAGE_LABELS.get(stage, stage)
                self._pending.append(
                    "[%s] %s (row %d) | %s | run elapsed %s"
                    % (stage, path, entry.row, label, self._run_elapsed(now)))

    def finish(self, path, row, result):
        now = self._clock()
        with self._lock:
            entry = self._entries.get(path)
            if entry is None or entry.result is not None:
                return
            entry.result = result
            entry.finished_at = now
            if entry.device_started is None:
                entry.device_started = (entry.probe_done or entry.stage_started
                                        or self._run_started)
            duration = _format_duration(
                (entry.finished_at - entry.device_started)
                if entry.device_started is not None else None)
            label = STAGE_LABELS.get(result, result)
            finished = sum(1 for item in self._entries.values() if item.result is not None)
            if not self._live:
                self._pending.append(
                    "[%s] %s (row %d) | %s | duration %s | finished %d/%d"
                    % (result, path, entry.row, label, duration, finished, self._total))

    def set_stopping(self):
        with self._lock:
            if self._stopping:
                return
            self._stopping = True
            if not self._live:
                self._pending.append(
                    "Stopping; waiting for active commands to exit.")

    # -- live refresh (collector thread) -------------------------------------

    def tick(self, force=False):
        if not self._live:
            self._drain()
            return
        now = self._clock()
        # Resolve the current geometry before formatting so a resize can never
        # produce lines wider than the terminal (which would wrap and desync
        # the cursor-up arithmetic).
        if self._size is None:
            columns, lines_count = _terminal_size(self._stream)
            with self._lock:
                if not (self._live and self._started):
                    return
                if not self._geometry_fits(columns, lines_count):
                    disable = True
                else:
                    disable = False
                    self._width = columns
            if disable:
                self._disable_live()
                return
        with self._lock:
            if not (self._live and self._started):
                return
            if not force and now - self._last_render < self._interval:
                return
            self._last_render = now
            lines = self._build_lines(now)
        with _PRINT_LOCK:
            ok = self._write_frame(lines)
        if not ok:
            with self._lock:
                self._live = False
                self._rendered = 0

    def _take_pending(self):
        with self._lock:
            pending = self._pending
            self._pending = []
        return pending

    def _write_pending_locked(self, pending):
        for line in pending:
            self._write(line + "\n")

    def _drain(self):
        pending = self._take_pending()
        if pending:
            with _PRINT_LOCK:
                self._write_pending_locked(pending)

    def _geometry_fits(self, columns, lines):
        return (columns >= _MIN_LIVE_WIDTH and lines >= _MIN_LIVE_HEIGHT
                and self._total <= min(_MAX_LIVE_ROWS, lines - 2))

    def _disable_live(self):
        """Leave live mode without cursor-up based on obsolete geometry."""
        with self._lock:
            if not self._live:
                return
            self._live = False
            rendered = self._rendered
            self._rendered = 0
        if rendered:
            with _PRINT_LOCK:
                self._write("\n")

    def _run_elapsed(self, now):
        if self._run_started is None:
            return "--:--"
        return _format_duration(now - self._run_started)

    def _build_lines(self, now):
        usable = max(1, self._width - 1)  # keep the last column free of auto-wrap
        lines = [_truncate(_sanitize_text(self._header_text(now)), usable)]
        state_width, row_width, duration_width = 24, 8, 8
        path_width = max(8, usable - (2 + row_width + 1 + state_width
                                      + 1 + 1 + duration_width))
        for path in self._order:
            entry = self._entries[path]
            label, duration = self._entry_state(entry, now)
            label = _truncate(_sanitize_text(label), state_width)
            path_text = _shorten_path(_sanitize_text(entry.path), path_width)
            row_text = ("row %d" % entry.row).ljust(row_width)
            colored = self._colorize(entry, label)
            padding = " " * max(0, state_width - len(label))
            lines.append("  %s %s%s %s %s" % (
                row_text, colored, padding, path_text.ljust(path_width),
                duration.rjust(duration_width)))
        return lines

    def _header_text(self, now):
        parts = [self._phase]
        if self._stopping:
            parts.append("stopping")
        if self._phase == "probing":
            done = probing = errors = 0
            for path in self._order:
                state = self._entries[path].probe_state
                if state in ("response", "error"):
                    done += 1
                elif state == "probing":
                    probing += 1
                if state == "error":
                    errors += 1
            parts += ["responded %d/%d" % (done, self._total),
                      "probing %d" % probing,
                      "waiting %d" % (self._total - done - probing),
                      "errors %d" % errors]
        else:
            finished = passed = failed = cancelled = unsaved = active = 0
            for path in self._order:
                entry = self._entries[path]
                if entry.result is not None:
                    finished += 1
                    if entry.result == RESULT_PASS:
                        passed += 1
                    elif entry.result == RESULT_FAIL:
                        failed += 1
                    elif entry.result == RESULT_CANCELLED:
                        cancelled += 1
                    elif entry.result == RESULT_NOT_SAVED:
                        unsaved += 1
                elif entry.stage is not None:
                    active += 1
            parts += ["done %d/%d" % (finished, self._total),
                      "pass %d" % passed, "fail %d" % failed,
                      "cancel %d" % cancelled, "unsaved %d" % unsaved,
                      "active %d" % active]
        parts.append(self._run_elapsed(now))
        return " | ".join(parts)

    def _entry_state(self, entry, now):
        if self._phase == "probing":
            label = _PROBE_STATE_LABELS[entry.probe_state]
            if entry.probe_state == "probing" and entry.probe_started is not None:
                duration = _format_duration(now - entry.probe_started)
            elif entry.probe_done is not None and entry.probe_started is not None:
                duration = _format_duration(entry.probe_done - entry.probe_started)
            else:
                duration = "--:--"
            return label, duration
        if entry.result is not None:
            label = STAGE_LABELS.get(entry.result, entry.result)
            duration = _format_duration(
                (entry.finished_at - entry.device_started)
                if entry.finished_at is not None and entry.device_started is not None
                else None)
            return label, duration
        if entry.stage is not None and entry.stage_started is not None:
            return STAGE_LABELS.get(entry.stage, entry.stage), _format_duration(
                now - entry.stage_started)
        if entry.device_started is not None:
            return "queued", _format_duration(now - entry.device_started)
        return "queued", "--:--"

    def _colorize(self, entry, text):
        if not self._use_color:
            return text
        if entry.result == RESULT_PASS:
            return "\033[32m%s\033[0m" % text
        if entry.result in (RESULT_FAIL, RESULT_NOT_SAVED):
            return "\033[31m%s\033[0m" % text
        return text

    def _write_frame(self, lines):
        """Write a frame in place; caller must hold _PRINT_LOCK.

        Returns False when the stream is unusable so the caller can fall back
        without letting display I/O abort the batch.
        """
        body = "\n".join("\x1b[K" + line for line in lines)
        prefix = ""
        if self._rendered:
            up = self._rendered - 1
            prefix = ("\x1b[%dA" % up if up else "") + "\r"
        try:
            self._stream.write(prefix + body)
            self._stream.flush()
        except (OSError, ValueError):
            return False
        self._rendered = len(lines)
        return True

    def _write(self, text):
        try:
            self._stream.write(text)
            self._stream.flush()
        except (OSError, ValueError):
            pass

    def _print_plain(self, text):
        with _PRINT_LOCK:
            self._write(text + "\n")

    def _print_plain_summary(self, phase):
        if phase == "probing":
            done = errors = 0
            for path in self._order:
                state = self._entries[path].probe_state
                if state in ("response", "error"):
                    done += 1
                if state == "error":
                    errors += 1
            self._print_plain("Probing complete: responded %d/%d | errors %d"
                              % (done, self._total, errors))
        elif phase == "flashing":
            counts = {RESULT_PASS: 0, RESULT_FAIL: 0,
                      RESULT_CANCELLED: 0, RESULT_NOT_SAVED: 0}
            finished = 0
            for path in self._order:
                result = self._entries[path].result
                if result is not None:
                    finished += 1
                    if result in counts:
                        counts[result] += 1
            self._print_plain(
                "Flashing complete: finished %d/%d | pass %d | fail %d | "
                "cancelled %d | unsaved %d"
                % (finished, self._total, counts[RESULT_PASS], counts[RESULT_FAIL],
                   counts[RESULT_CANCELLED], counts[RESULT_NOT_SAVED]))


def _color(text, status):
    if not _color_enabled(sys.stdout):
        return text
    if status == RESULT_PASS:
        return "\033[32m%s\033[0m" % text
    if status in (RESULT_FAIL, RESULT_NOT_SAVED):
        return "\033[31m%s\033[0m" % text
    return text


def print_candidates(candidates):
    print("  %-30s %-9s %-16s %-16s %s" % (
        "DEVICE", "VID:PID", "USB SERIAL", "USB LOCATION", "DESCRIPTION"))
    for candidate in candidates:
        print("  %-30s %-9s %-16s %-16s %s" % (
            candidate.path, candidate.vid_pid or "-", candidate.serial_number or "-",
            candidate.location or "-", candidate.description or "-"))


def print_preview(assignments, retry, dry_run):
    title = ("Provisional device -> credential row mapping (dry-run)" if dry_run
             else "Device -> credential row mapping")
    print("\n%s:" % title)
    print("  %-30s %-9s %-16s %-16s %4s  %-24s %s" % (
        "DEVICE", "VID:PID", "USB SERIAL", "USB LOCATION", "ROW", "UUID", "DESCRIPTION"))
    for assignment in assignments:
        candidate = assignment.candidate
        print("  %-30s %-9s %-16s %-16s %4d  %-24s %s" % (
            candidate.path, candidate.vid_pid or "-", candidate.serial_number or "-",
            candidate.location or "-", assignment.row.index,
            assignment.row.uuid, candidate.description or "-"))
    if retry:
        print("  (retry: reusing the recorded row and chip identity)")
    print("\nWARNING: flashing erases/programs the listed images and replaces the "
          "entire NVS partition, including WiFi, settings and activation. Probing "
          "may reset devices. Chip compatibility does not verify the board variant; "
          "check the build/board and USB cables, and unplug unrelated devices "
          "or select a subset with --device.")


def print_report(assignments, results):
    print("\n%-30s %4s %-16s %-12s %s" % (
        "DEVICE", "ROW", "RESULT", "STAGE", "DETAIL"))
    failed = []
    for assignment in assignments:
        result = results.get(assignment.candidate.path)
        if result is None:
            continue
        print("%-30s %4d %-16s %-12s %s" % (
            result.path, result.row, _color(result.result, result.result),
            result.stage, result.detail or ""))
        if result.result in (RESULT_FAIL, RESULT_NOT_SAVED):
            failed.append((assignment, result))
    if failed:
        print("\nFAILED DEVICES:")
        for assignment, result in failed:
            print("  %s - %s: %s" % (result.path, result.stage,
                                     result.detail or result.result))
            candidate = assignment.candidate
            metadata = []
            if candidate.location:
                metadata.append("USB location: %s" % candidate.location)
            if candidate.serial_number:
                metadata.append("USB serial: %s" % candidate.serial_number)
            if assignment.mac:
                metadata.append("chip: %s" % assignment.mac)
            if metadata:
                print("    " + "; ".join(metadata))


# ---------------------------------------------------------------------------
# Commands
# ---------------------------------------------------------------------------

def cmd_list_ports(args):
    candidates, excluded = discover_candidates()
    if candidates:
        print("USB serial candidates (%d):" % len(candidates))
        print_candidates(candidates)
    else:
        print("No USB serial candidates found.")
    if excluded:
        print("\nExcluded ports:")
        for device, reason in excluded:
            print("  %s - %s" % (device, reason))
    if os.environ.get("ESPPORT"):
        print("\nnote: $ESPPORT is not used for batch discovery.")
    return 0


def cmd_dry_run(args, cancel_event):
    args.xlsx = os.path.realpath(args.xlsx)
    _require_pid(args)
    _validate_idf_tools()
    manifest = Manifest(args.build_dir).load()
    if args.retry_uuid:
        _validate_retry_structure(args)
    candidates = _resolve_candidates(args)
    workbook = Workbook(args.xlsx, args.sheet).load(for_write=False)
    assignments = plan_assignments(args, workbook, candidates)
    print_preview(assignments, bool(args.retry_uuid), dry_run=True)
    print("\ndry-run: no ports opened, no probing, no workbook changes.")
    return 0


def cmd_flash(args, cancel_event):
    args.xlsx = os.path.realpath(args.xlsx)
    _require_pid(args)
    _validate_idf_tools()
    if args.retry_uuid:
        _validate_retry_structure(args)
    manifest = Manifest(args.build_dir).load()
    candidates = _resolve_candidates(args)

    lock = WorkbookLock(args.xlsx)
    lock.acquire()
    journal = None
    coordinator = None
    try:
        workbook = Workbook(args.xlsx, args.sheet).load(for_write=True)
        assignments = plan_assignments(args, workbook, candidates)

        print("\nDiscovered candidate devices (%d):" % len(candidates))
        print_candidates(candidates)
        print_preview(assignments, bool(args.retry_uuid), dry_run=False)
        if not args.yes:
            _confirm(cancel_event)

        check_cancelled(cancel_event, "preflight")
        backup = write_backup(args.xlsx)
        check_cancelled(cancel_event, "preflight")
        journal = Journal(args.xlsx)
        check_cancelled(cancel_event, "preflight")
        verify_workbook_directory_writable(args.xlsx)

        check_cancelled(cancel_event, "probe")
        progress = ProgressDisplay(assignments)
        progress.start("probing")
        try:
            probe_results = probe_barrier(candidates, args, manifest, cancel_event,
                                          progress)
        finally:
            progress.stop()
        if cancel_event.is_set():
            print("\nCANCELLED before the preflight barrier completed.", file=sys.stderr)
            return 130
        validate_probes(probe_results, manifest, workbook, args)
        for assignment in assignments:
            assignment.mac = probe_results[assignment.candidate.path].mac
        manifest.verify_fingerprint()
        # Freeze the probed snapshot before any credential is reserved.
        verify_candidates_frozen(candidates)
        if cancel_event.is_set():
            print("\nCANCELLED before any credential was reserved.", file=sys.stderr)
            return 130

        run_id = _new_run_id()
        coordinator = Coordinator(workbook, journal, run_id, cancel_event, progress)
        coordinator.register_secrets(assignments)
        print("\nReserving %d credential row(s) (backup: %s)..." % (len(assignments), backup))
        coordinator.reserve(assignments, args.pid)

        jobs = args.jobs or len(assignments)
        print("Flashing %d device(s) with up to %d concurrent job(s)..." % (
            len(assignments), max(1, min(jobs, len(assignments)))))
        progress.start("flashing")
        summarize = False
        try:
            results = run_workers(assignments, manifest, coordinator, jobs,
                                  args.baud, args.timeout)
            # Close the journal before publishing any phase summary: a late
            # close failure must surface as RESULT NOT SAVED, never behind an
            # all-success progress footer.
            try:
                journal.close()
                summarize = True
            except Exception as exc:
                raise coordinator.record_persistence_failure(exc) from exc
            finally:
                journal = None
        finally:
            progress.stop(summarize=summarize)
        for result in results.values():
            if result.result == RESULT_NOT_SAVED:
                result.detail = str(coordinator.record_persistence_failure(
                    PersistenceError(result.detail or RESULT_NOT_SAVED)))
        print_report(assignments, results)

        if coordinator.failure is not None:
            print("\nRESULT NOT SAVED: %s" % coordinator.failure, file=sys.stderr)
            return 1
        if cancel_event.is_set():
            print("\nCANCELLED: some rows remain in_progress and are not reusable "
                  "without explicit retry.", file=sys.stderr)
            return 130
        if any(result.result != RESULT_PASS for result in results.values()):
            return 1
        print("\nAll %d device(s) verified and saved as used." % len(results))
        return 0
    except PersistenceError as exc:
        cancel_event.set()
        print("\nRESULT NOT SAVED: %s" % exc, file=sys.stderr)
        return 1
    finally:
        try:
            if journal is not None:
                try:
                    journal.close()
                except Exception as exc:
                    error = (coordinator.record_persistence_failure(exc)
                             if coordinator is not None else exc)
                    print("\nRESULT NOT SAVED: %s" % error, file=sys.stderr)
                    return 1
        finally:
            lock.release()


def _confirm(cancel_event):
    check_cancelled(cancel_event, "preflight")
    if not sys.stdin.isatty():
        raise BatchError("preflight", "refusing to flash non-interactively",
                         "pass --yes to authorize unattended flashing (prefer --device)")
    print("Type 'yes' to continue: ", end="", flush=True)
    while True:
        check_cancelled(cancel_event, "preflight")
        ready, _, _ = select.select([sys.stdin], [], [], 0.1)
        check_cancelled(cancel_event, "preflight")
        if ready:
            answer = sys.stdin.readline().strip().lower()
            check_cancelled(cancel_event, "preflight")
            break
    if answer != "yes":
        raise BatchError("preflight", "aborted by operator")


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def _utc_now():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _new_run_id():
    return time.strftime("%Y%m%dT%H%M%SZ", time.gmtime()) + "-" + uuid.uuid4().hex[:8]


def _install_signal_handlers(cancel_event):
    def handler(_signum, _frame):
        cancel_event.set()
    for signum in (signal.SIGINT, signal.SIGTERM):
        try:
            signal.signal(signum, handler)
        except (ValueError, OSError):  # pragma: no cover - non-main thread
            pass


def parse_args(argv):
    parser = argparse.ArgumentParser(
        prog="batch_flash.py", description=__doc__,
        formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--pid", help="Tuya product key (required to flash/dry-run/retry)")
    parser.add_argument("--list-ports", action="store_true",
                        help="enumerate USB serial candidates and exit")
    parser.add_argument("--dry-run", action="store_true",
                        help="discovery + local validation + row preview; no device access")
    parser.add_argument("--yes", action="store_true",
                        help="authorize the displayed candidate set without prompting")
    parser.add_argument("--device", action="append", default=[], metavar="PORT",
                        help="restrict to this port path (repeatable; no LABEL= syntax)")
    parser.add_argument("--xlsx", default=None,
                        help="credential workbook (default: <project>/auth-info.xlsx)")
    parser.add_argument("--sheet", default=None,
                        help="worksheet name (default: the sheet with uuid/key headers)")
    parser.add_argument("--project-dir", default=None,
                        help="project root for default paths (default: parent of scripts/)")
    parser.add_argument("--build-dir", default=None,
                        help="build dir with flasher_args.json (default: <project>/build)")
    parser.add_argument("--baud", type=int, default=None,
                        help="serial baud rate for esptool (default: esptool default)")
    parser.add_argument("--jobs", type=int, default=None,
                        help="max concurrent device jobs (default: number of devices)")
    parser.add_argument("--timeout", type=float, default=DEFAULT_TIMEOUT,
                        help="per-command timeout seconds (default: %(default)s)")
    parser.add_argument("--retry-uuid", default=None, metavar="UUID",
                        help="retry one failed/interrupted row; requires one --device")
    parser.add_argument("-p", "--port", default=os.environ.get("ESPPORT"),
                        help=argparse.SUPPRESS)
    args = parser.parse_args(argv)

    project_dir = (args.project_dir or os.environ.get("IDF_PROJECT_DIR")
                   or os.path.dirname(_HERE))
    args.project_dir = os.path.abspath(project_dir)
    build_dir = (args.build_dir or os.environ.get("IDF_BUILD_DIR")
                 or os.path.join(args.project_dir, "build"))
    args.build_dir = os.path.abspath(build_dir)
    args.xlsx = os.path.realpath(args.xlsx or os.path.join(args.project_dir,
                                                          "auth-info.xlsx"))
    if args.baud is not None and args.baud <= 0:
        parser.error("--baud must be positive")
    if args.jobs is not None and args.jobs < 1:
        parser.error("--jobs must be positive")
    if not math.isfinite(args.timeout) or args.timeout <= 0:
        parser.error("--timeout must be finite and positive")
    return args


def main(argv=None):
    args = parse_args(argv)
    cancel_event = threading.Event()
    _install_signal_handlers(cancel_event)
    try:
        if args.port:
            raise BatchError(
                "preflight", "-p/--port is not valid for batch flashing",
                "omit it to use USB discovery, use --device for an explicit "
                "subset, and clear ESPPORT if it sets the global port")
        if args.list_ports:
            return cmd_list_ports(args)
        if args.dry_run:
            return cmd_dry_run(args, cancel_event)
        return cmd_flash(args, cancel_event)
    except BatchError as exc:
        print("error: %s" % exc, file=sys.stderr)
        return exc.exit_code
    except (CommandCancelled, KeyboardInterrupt):
        cancel_event.set()
        print("\ncancelled", file=sys.stderr)
        return 130


if __name__ == "__main__":
    sys.exit(main())
